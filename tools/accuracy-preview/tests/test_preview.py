"""Synthetic-only preview regression, causality, and component integrity checks."""

import json
import math
from datetime import datetime
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from strath_accuracy_preview import cli, core

from strathmark.v3.assessors.ml import PITCalibrator, build_positive_distribution
from strathmark.v3.contracts.canonical import canonical_digest
from strathmark.v3.contracts.evidence import TargetContext
from strathmark.v3.contracts.forecasts import PositiveTimeDistribution, QuantilePoint
from strathmark.v3.factory.ml_training import FEATURE_NAMES
from strathmark.v3.factory.workbook_history import load_workbook_history

RAW = tuple(math.log(v) for v in (10, 20, 40, 60, 80, 100, 120))
IDENTITY = PITCalibrator(
    "calibration", (("0", "0"), ("1", "1")), "a" * 64, "strathmark-v3-ml-pit-calibrator-v2", "0.2"
)
CONTEXT = TargetContext("underhand", 300, "synthetic-pine", "strathex:v1", "strathex:v1", ())


@pytest.fixture
def calibration():
    return {
        "supported_scale": 1.25,
        "supported_rows": 24,
        "supported_groups": 8,
        "cold_rows": 24,
        "cold_groups": 8,
        "cold_formula": IDENTITY.to_dict(),
        "cold_ml": IDENTITY.to_dict(),
    }


@pytest.fixture
def artifact(tmp_path, monkeypatch, calibration):
    import strathmark.v3.factory.ml_artifacts as artifacts
    import strathmark.v3.linux_forecasts as native
    import strathmark.v3.runtime_identity as runtime

    root = tmp_path / "candidate"
    (root / "components/ml-bundle").mkdir(parents=True)
    (root / "components/ml-bundle/manifest.json").write_text('{"synthetic":true}')
    (root / "components/formula_manifest.json").write_text('{"synthetic":true}')
    core.write_new_json(root / "calibration.json", calibration)
    bundle = SimpleNamespace(
        digest=core.sha(root / "components/ml-bundle/manifest.json"),
        specialist_models={},
        feature_names=FEATURE_NAMES,
        calibrator=IDENTITY,
        normalize_features=lambda f: (f, ()),
        universal_model=object(),
    )
    monkeypatch.setattr(artifacts, "load_ml_bundle", lambda *a, **k: bundle)
    monkeypatch.setattr(
        native, "load_formula_manifest", lambda *a: SimpleNamespace(digest="b" * 64)
    )
    monkeypatch.setattr(runtime, "implementation_digest", lambda: "c" * 64)
    manifest = {
        "schema_version": core.SCHEMA,
        "purpose": core.PURPOSE,
        "eligible_for_official_use": False,
        "baseline_bundle_digest": bundle.digest,
        "formula_digest": "b" * 64,
        "implementation_digest": "c" * 64,
        "files": {
            p.relative_to(root).as_posix(): core.sha(p) for p in root.rglob("*") if p.is_file()
        },
        "provenance": {"synthetic_only": True},
    }
    core.write_new_json(
        root / "manifest.json", {**manifest, "artifact_digest": canonical_digest(manifest)}
    )
    return root, bundle


def workbook(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Competitor"
    sheet.append(["CompetitorID"])
    sheet.append(["SYN001"])
    wood = book.create_sheet("Wood")
    wood.append(["speciesID", "spec_gravity"])
    wood.append(["synthetic-pine", 0.6])
    results = book.create_sheet("Results")
    results.append(
        ["CompetitorID", "Event", "Time (seconds)", "Size (mm)", "Species Code", "Date (optional)"]
    )
    for date, seconds in (
        (datetime(2020, 1, 1), 40),
        (datetime(2020, 3, 1), 500),
        (datetime(2020, 3, 2), 599),
    ):
        results.append(["SYN001", "UH", seconds, 300, "synthetic-pine", date])
    book.save(path)
    book.close()


def test_supported_scaling_preserves_originals_and_intervals(calibration):
    dist = build_positive_distribution(RAW, IDENTITY)
    original = dist.to_dict()
    formula, ml, pool = core.apply_calibration(dist, dist, RAW, 3, calibration)
    assert dist.to_dict() == original
    assert formula.median_ms == ml.median_ms == pool.median_ms == 75000
    assert all(1 <= q.time_ms <= 600000 for q in formula.quantiles)
    assert [q.time_ms for q in formula.quantiles] == sorted(q.time_ms for q in formula.quantiles)


def test_cold_maps_use_raw_quantiles_without_supported_scale(calibration):
    dist = build_positive_distribution(RAW, IDENTITY)
    changed = PITCalibrator(
        "calibration",
        (("0", "0"), ("0.9", "0.5"), ("1", "1")),
        "d" * 64,
        "strathmark-v3-ml-pit-calibrator-v2",
        "0.5",
    )
    calibration["cold_ml"] = calibration["cold_formula"] = changed.to_dict()
    coarse = PositiveTimeDistribution(
        tuple(
            QuantilePoint(p, t)
            for p, t in (
                ("0.05", 10000),
                ("0.25", 40000),
                ("0.5", 60000),
                ("0.75", 80000),
                ("0.95", 120000),
            )
        )
    )
    formula, ml, pool = core.apply_calibration(coarse, dist, RAW, 0, calibration)
    assert ml.median_ms == 100000
    assert formula.median_ms == 110000
    assert 100000 <= pool.median_ms <= 110000


@pytest.mark.parametrize("scale", [0, -1, True, float("nan"), float("inf")])
def test_invalid_scale_refused(calibration, scale):
    calibration["supported_scale"] = scale
    with pytest.raises(ValueError):
        core.validate_calibration(calibration)


@pytest.mark.parametrize(
    "key,value", [("cold_rows", 19), ("cold_groups", 4), ("supported_groups", True)]
)
def test_insufficient_actual_support_refused(calibration, key, value):
    calibration[key] = value
    with pytest.raises(ValueError):
        core.validate_calibration(calibration)


def test_loader_accepts_only_unchanged_components(artifact):
    root, bundle = artifact
    assert core.load_candidate(root).bundle is bundle
    (root / "calibration.json").write_text("{}")
    with pytest.raises(ValueError, match="component bytes"):
        core.load_candidate(root)


def test_loader_refuses_wrong_python_abi(artifact, monkeypatch):
    root, _ = artifact
    monkeypatch.setattr(core, "sys", SimpleNamespace(version_info=(3, 12)))
    with pytest.raises(ValueError, match="Python 3.13"):
        core.load_candidate(root)


def test_cold_floor_cannot_be_narrowed(artifact):
    root, _ = artifact
    calibration = json.loads((root / "calibration.json").read_text())
    calibration["cold_ml"]["interval_log_radius"] = "0.1"
    (root / "calibration.json").write_text(json.dumps(calibration))
    manifest = json.loads((root / "manifest.json").read_text())
    manifest["files"]["calibration.json"] = core.sha(root / "calibration.json")
    body = {k: v for k, v in manifest.items() if k != "artifact_digest"}
    manifest["artifact_digest"] = canonical_digest(body)
    (root / "manifest.json").write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match="narrower"):
        core.load_candidate(root)


@pytest.mark.parametrize("mutation", ["authority", "unknown", "extra", "symlink", "runtime"])
def test_manifest_and_file_tampering_refused(artifact, mutation, monkeypatch):
    root, _ = artifact
    manifest = json.loads((root / "manifest.json").read_text())
    if mutation == "authority":
        manifest["eligible_for_official_use"] = True
    elif mutation == "unknown":
        manifest["issued_mark"] = True
    elif mutation == "extra":
        (root / "components/extra.json").write_text("{}")
    elif mutation == "symlink":
        p = root / "components/formula_manifest.json"
        original = root.parent / "formula.json"
        p.rename(original)
        p.symlink_to(original)
    else:
        import strathmark.v3.runtime_identity as runtime

        monkeypatch.setattr(runtime, "implementation_digest", lambda: "e" * 64)
    (root / "manifest.json").write_text(json.dumps(manifest))
    with pytest.raises(ValueError):
        core.load_candidate(root)


def test_snapshot_is_verified_readonly_and_refuses_overwrite(tmp_path):
    original, copied = tmp_path / "source.xlsx", tmp_path / "copy.xlsx"
    workbook(original)
    before = core.sha(original)
    receipt = core.snapshot(original, copied)
    assert core.sha(original) == core.sha(copied) == receipt["sha256"] == before
    core.readonly_workbook(copied)
    with pytest.raises(ValueError, match="snapshot"):
        core.readonly_workbook(original)
    with pytest.raises(FileExistsError):
        core.snapshot(original, copied)
    assert core.sha(copied) == before


@pytest.mark.parametrize(
    "cutoff,prior_count", [("2020-01-01T00:00:00.000Z", 0), ("2020-03-01T00:00:00.000Z", 1)]
)
@pytest.mark.parametrize("exact_native", [True, False])
def test_forecast_matches_native_features_excludes_target_day_and_never_issues(
    artifact, tmp_path, monkeypatch, cutoff, prior_count, exact_native
):
    import strathmark.v3.factory.ml_training as training
    import strathmark.v3.linux_forecasts as native

    root, _ = artifact
    source, copied = tmp_path / "source.xlsx", tmp_path / "copy.xlsx"
    workbook(source)
    core.snapshot(source, copied)
    monkeypatch.setattr(training, "predict_model_log_quantiles", lambda *a: RAW)
    dist = build_positive_distribution(RAW, IDENTITY)
    calls = []

    def calculate(**kwargs):
        calls.append(kwargs)
        assert "field_id" not in kwargs
        history = load_workbook_history(
            copied, cutoff_at_utc=kwargs["round_snapshot"]["cutoff_at_utc"]
        )
        assert len(history.observations) == prior_count
        identifier = dict(history.competitor_ids)["SYN001"]
        return {
            "issued_mark": False,
            "forecasts": [
                {
                    "local_competitor_id": "SYN001",
                    "competitor_id": identifier,
                    "formula": {"forecast": {"distribution": dist.to_dict()}},
                    "ml": {
                        "forecast": {
                            "distribution": (
                                dist
                                if exact_native
                                else build_positive_distribution(
                                    tuple(v + 0.1 for v in RAW), IDENTITY
                                )
                            ).to_dict()
                        }
                    },
                    "predicted_time_ms": dist.median_ms,
                }
            ],
        }

    monkeypatch.setattr(native, "calculate", calculate)
    if not exact_native:
        with pytest.raises(ValueError, match="no approximate substitution"):
            core.forecast(core.load_candidate(root), copied, cutoff, ["SYN001"], CONTEXT.to_dict())
        return
    result = core.forecast(core.load_candidate(root), copied, cutoff, ["SYN001"], CONTEXT.to_dict())
    assert result["issued_mark"] is result["eligible_for_official_use"] is False
    assert result["forecasts"][0]["same_event_history_depth"] == prior_count
    assert result["forecasts"][0]["candidate_time_ms"] == (75000 if prior_count else 60000)
    assert calls[0]["round_snapshot"]["live_results"] == []
    assert "marks" not in json.dumps(result)


def test_menu_builds_closed_target_context(artifact, tmp_path, monkeypatch, capsys):
    root, _ = artifact
    source, copied = tmp_path / "source.xlsx", tmp_path / "copy.xlsx"
    workbook(source)
    core.snapshot(source, copied)
    answers = iter(["2", "2020-03-01T00:00:00.000Z", "SYN001", "UH", "synthetic-pine", "300", "0"])
    monkeypatch.setattr("builtins.input", lambda _: next(answers))
    captured = []

    def predict(candidate, path, cutoff, ids, context):
        captured.append(TargetContext.from_dict(context))
        return {"forecasts": []}

    monkeypatch.setattr(cli, "forecast", predict)
    cli.menu(core.load_candidate(root), copied, None)
    assert len(captured) == 1
    assert captured[0].event_code == "underhand"
    assert "Preview refused" not in capsys.readouterr().out


def test_cli_rejects_field_request_and_preserves_existing_output(artifact, tmp_path, capsys):
    root, _ = artifact
    request = tmp_path / "request.json"
    request.write_text(
        json.dumps(
            {
                "workbook": "unused.xlsx",
                "cutoff_at_utc": "2020-03-01T00:00:00.000Z",
                "competitor_ids": ["SYN001"],
                "target_context": CONTEXT.to_dict(),
                "field_id": "field:synthetic",
            }
        )
    )
    assert cli.main(["forecast", "--candidate", str(root), "--request", str(request)]) == 2
    assert "cutting-time fields only" in capsys.readouterr().err
    path = tmp_path / "existing.json"
    path.write_text("preserved")
    with pytest.raises(FileExistsError):
        core.write_new_json(path, {})
    assert path.read_text() == "preserved"
