"""Legacy candidate ingress uses only synthetic workbook facts and never writes them."""

from datetime import datetime
from hashlib import sha256

import pytest
from openpyxl import Workbook

from strathmark.v3.factory.workbook_history import load_workbook_history


def workbook_history_fixture(path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    competitors = workbook.create_sheet("Competitor")
    competitors.append(["CompetitorID", "Name"])
    competitors.append(["SYN001", "Synthetic competitor"])
    wood = workbook.create_sheet("Wood")
    wood.append(["speciesID", "spec_gravity"])
    wood.append(["S01", 0.5])
    results = workbook.create_sheet("Results")
    results.append(
        [
            "CompetitorID",
            "Event",
            "Time (seconds)",
            "Size (mm)",
            "Species Code",
            "Date (optional)",
            "Notes (Competition, special circumstances, etc.)",
        ]
    )
    results.append(["SYN001", "UH", 28, 300, "S01", datetime(2022, 1, 2), "Annual synthetic event"])
    results.append(["SYN001", "SB", 34, 300, "S01", datetime(2023, 1, 2), "Annual synthetic event"])
    results.append(["SYN001", "UH", 29, 300, "S01", None, None])
    results.append(["SYN001", "UH", 30, 300, "S01", datetime(2024, 1, 2), None])
    workbook.save(path)
    workbook.close()


def test_exclusive_cutoff_dated_rows_canonical_units_and_no_live_issue_facts(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    workbook_history_fixture(path)
    before = sha256(path.read_bytes()).hexdigest()
    history = load_workbook_history(path, cutoff_at_utc="2024-01-02T00:00:00.000Z")
    assert history.source_sha256 == before == sha256(path.read_bytes()).hexdigest()
    assert dict(history.excluded) == {"undated": 1, "at_or_after_cutoff": 1}
    assert [item.context.event_code for item in history.observations] == [
        "underhand",
        "standing_block",
    ]
    assert [item.result.raw_time_ms for item in history.observations] == [28000, 34000]
    assert len({item.tournament_id for item in history.observations}) == 2
    for item in history.observations:
        assert item.context.properties[0].value == "500"
        assert item.context.properties[0].unit == "kg_m3"
        assert item.completion_clock_ms is None
        assert item.placing is None
        assert item.gap_ms is None
    assert "Synthetic competitor" not in str(history)


def test_naive_cutoff_is_rejected(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    workbook_history_fixture(path)
    with pytest.raises(ValueError, match="timezone"):
        load_workbook_history(path, cutoff_at_utc="2024-01-02T00:00:00")


def test_malformed_material_is_excluded_without_losing_valid_rows(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "synthetic.xlsx"
    workbook_history_fixture(path)
    workbook = load_workbook(path)
    workbook["Results"].append(["SYN001", "UH", 30, 300, None, datetime(2023, 2, 2), None])
    workbook.save(path)
    workbook.close()
    history = load_workbook_history(path, cutoff_at_utc="2025-01-01T00:00:00.000Z")
    assert len(history.observations) == 3
    assert dict(history.excluded)["invalid_material_context"] == 1


def test_packaged_formula_manifest_matches_reviewed_source():
    from pathlib import Path

    import strathmark.v3.linux_candidate as candidate

    root = Path(__file__).resolve().parents[3]
    assert (Path(candidate.__file__).parent / "contracts/formula_manifest.json").read_bytes() == (
        root / "benchmarks/v3/formula_manifest.json"
    ).read_bytes()
