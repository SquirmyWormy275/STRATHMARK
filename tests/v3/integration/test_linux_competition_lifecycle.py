from __future__ import annotations

import math
import sys
from datetime import datetime

import pytest

pytestmark = pytest.mark.skipif(
    sys.version_info[:2] != (3, 13), reason="Linux V3 execution requires Python 3.13"
)

from strathmark.v3.linux_lifecycle import CONTRACT_DIGEST, LinuxCompetitionRuntime
from strathmark.v3.linux_lifecycle_store import LinuxLifecycleError, LinuxLifecycleStore, digest

CUTOFF = "2026-10-02T00:00:00.000Z"


@pytest.fixture(scope="module")
def trained_synthetic_bundle(tmp_path_factory):
    catboost = pytest.importorskip("catboost")
    from strathmark.v3.assessors.ml import PITCalibrator, SpecialistGate
    from strathmark.v3.factory.ml_artifacts import (
        BUNDLE_METADATA_SCHEMA,
        DEPENDENCY_SCHEMA,
        FEATURE_SCHEMA,
        VOCABULARY_SCHEMA,
        export_catboost_json,
        write_ml_bundle,
    )
    from strathmark.v3.factory.ml_training import (
        CATEGORICAL_FEATURES,
        FEATURE_NAMES,
        CausalTrainingRow,
        _train_catboost_hierarchy,
    )

    root = tmp_path_factory.mktemp("synthetic-linux-trained-model")
    rows = []
    for index in range(40):
        seconds = 28 if index % 2 == 0 else 40
        features = {
            "event_family": "underhand",
            "species": "gum",
            "size_mm": 300,
            "density": 500.0,
            "density_missing": 0,
            "history_depth": index,
            "exact_history_depth": index,
            "history_log_median": math.log(seconds),
            "history_log_spread": 0.1,
            "history_missing": 0,
            "sequence_recency": 0,
            "history_log_trend": 0.0,
            "context_distance": 0.0,
            "eligible_tournament_sequence": index,
            "current_form_log_seconds": math.log(seconds),
        }
        rows.append(
            CausalTrainingRow(
                f"evidence:synthetic-{index}",
                f"competitor:synthetic-{index % 2}",
                f"tournament:synthetic-training-{index % 5}",
                "2025-01-02T00:00:00.000Z",
                index + 1,
                "underhand|300|gum",
                tuple((name, features[name]) for name in FEATURE_NAMES),
                str(math.log(seconds)),
                "a" * 64,
                index,
                "2025-01-01T00:00:00.000Z" if index else "0001-01-01T00:00:00.000Z",
                f"field:synthetic-{index}",
                "strathex:v1",
                "strathex:v1",
            )
        )
    universal, specialists, _eligibility = _train_catboost_hierarchy(
        tuple(rows), iterations=12, depth=2
    )
    assert specialists == {}
    calibrator = PITCalibrator.identity(source_digest="e" * 64)
    return write_ml_bundle(
        root / "ml-bundle",
        universal_model_json=export_catboost_json(universal, root / "universal.json"),
        specialist_model_json={},
        gate=SpecialistGate("0", (("log_history_depth", "0"), ("missing_fraction", "0"))),
        calibrator=calibrator,
        feature_schema={
            "schema_version": FEATURE_SCHEMA,
            "features": list(FEATURE_NAMES),
            "categorical": list(CATEGORICAL_FEATURES),
            "quantiles": ["0.05", "0.1", "0.25", "0.5", "0.75", "0.9", "0.95"],
        },
        category_vocabulary={
            "schema_version": VOCABULARY_SCHEMA,
            "values": {"event_family": ["__other__", "underhand"], "species": ["__other__", "gum"]},
        },
        dependency_lock={
            "schema_version": DEPENDENCY_SCHEMA,
            "catboost_version": catboost.__version__,
            "python_abi": "cp313",
        },
        bundle_metadata={
            "schema_version": BUNDLE_METADATA_SCHEMA,
            "code_revision": "synthetic-integration-test",
            "training_snapshot_digest": "1" * 64,
            "role_manifest_digest": "2" * 64,
            "gate_oof_digest": "3" * 64,
            "calibrator_source_digest": calibrator.source_digest,
            "taxonomy_version": "strathex:v1",
            "conversion_version": "strathex:v1",
        },
        bundle_version="ml:synthetic-linux-test",
    )


@pytest.fixture
def competition(tmp_path, trained_synthetic_bundle):
    from openpyxl import Workbook

    from strathmark.v3.factory.workbook_history import load_workbook_history

    workbook = Workbook()
    workbook.remove(workbook.active)
    competitors = workbook.create_sheet("Competitor")
    competitors.append(["CompetitorID", "Name"])
    competitors.append(["SYN001", "Synthetic One"])
    competitors.append(["SYN002", "Synthetic Two"])
    wood = workbook.create_sheet("Wood")
    wood.append(["speciesID", "spec_gravity"])
    wood.append(["gum", 0.5])
    results = workbook.create_sheet("Results")
    results.append(
        ["CompetitorID", "Event", "Time (seconds)", "Size (mm)", "Species Code", "Date (optional)"]
    )
    for year in range(2021, 2026):
        for competitor, seconds in (("SYN001", 28), ("SYN002", 40)):
            results.append([competitor, "UH", seconds, 300, "gum", datetime(year, 1, 1)])
    path = tmp_path / "synthetic.xlsx"
    workbook.save(path)
    workbook.close()
    root = tmp_path / "synthetic-installation"
    LinuxLifecycleStore.initialize(root)
    runtime = LinuxCompetitionRuntime(root=root, ml_bundle=trained_synthetic_bundle)
    context = {
        "scope_id": "tournament:synthetic-show",
        "selected_engine": "v3",
        "mode": "local",
        "contract_identity": CONTRACT_DIGEST,
        "source_identity": runtime.status()["source_identity"],
        "selected_by_actor_id": "synthetic-judge",
        "locked_at": CUTOFF,
    }
    target = load_workbook_history(path, cutoff_at_utc=CUTOFF).observations[0].context.to_dict()
    request = {
        "workbook": str(path),
        "cutoff_at_utc": CUTOFF,
        "round_id": "round:synthetic-heats",
        "round_ordinal": 1,
        "epoch_group_id": "round:synthetic-event",
        "predecessor_round_ids": [],
        "competitor_ids": ["SYN001", "SYN002"],
        "upstream_competitor_ids": ["competitor:synthetic-1", "competitor:synthetic-2"],
        "target_context": target,
        "field_id": "field:synthetic-heat-1",
        "field_kind": "handicap",
        "upstream_field_revision": 1,
        "stand_ids": ["stand:synthetic-1", "stand:synthetic-2"],
        "ceiling": 180,
    }
    return runtime, context, request, root


def invoke(runtime, context, operation, payload, *, key=None):
    return runtime.execute(
        operation,
        {
            "command_id": key or "test:" + digest({"operation": operation, "payload": payload}),
            "context": context,
            "payload": payload,
        },
    )


def approve_and_issue(runtime, context, receipt):
    page = invoke(runtime, context, "approval_page", {"offset": 0, "limit": 100})
    row = next(row for row in page["rows"] if row["receipt_id"] == receipt["receipt_id"])
    binding = {
        "field_id": row["field_id"],
        "receipt_id": row["receipt_id"],
        "receipt_digest": row["receipt_content_digest"],
        "receipt_revision": row["receipt_revision"],
        "upstream_field_revision": row["upstream_field_revision"],
        "row_digest": row["row_digest"],
        "call_order": row["call_order"],
    }
    action = "individual_accept" if row["classification"] == "red" else "degraded_batch_accept"
    decision = invoke(
        runtime,
        context,
        "approve",
        {
            "schema_version": "strathmark-v3-approval-decision-request-v1",
            "tournament_id": context["scope_id"],
            "snapshot_id": page["snapshot_id"],
            "action": action,
            "selected": [binding],
            "excluded": [],
            "actor_metadata": {
                "asserted_actor_id": "synthetic-judge",
                "trust_model": "local_os_user",
            },
            "reason_code": "synthetic_deliberate_review",
            "superseded_receipt_id": None,
            "decided_at_utc": "2026-10-02T21:00:00.000Z",
            "deadline_ms": 10000,
        },
    )
    assert decision["issued"] is False
    return invoke(
        runtime,
        context,
        "issue",
        {
            "schema_version": "strathmark-v3-issue-acknowledgment-request-v1",
            "upstream_issue_id": "issue:" + receipt["receipt_id"].split(":")[1],
            "receipt_bindings": [
                {"receipt_id": receipt["receipt_id"], "receipt_digest": receipt["receipt_digest"]}
            ],
            "issued_at_utc": "2026-10-02T21:00:01.000Z",
            "deadline_ms": 10000,
        },
    )


def settle(runtime, context, receipt, issue, *, results=None, observed_at=None):
    return invoke(
        runtime,
        context,
        "settle",
        {
            "schema_version": "strathmark-v3-settlement-request-v1",
            "receipt_id": receipt["receipt_id"],
            "issue_batch_id": issue["issue_batch_id"],
            "results": results
            or [
                {
                    "competitor_id": receipt["competitor_ids"][0],
                    "status": "completion",
                    "raw_time_ms": 22000,
                    "penalty_ms": None,
                    "source_revision": 1,
                    "official_placing": 1,
                },
                {
                    "competitor_id": receipt["competitor_ids"][1],
                    "status": "completion",
                    "raw_time_ms": 50000,
                    "penalty_ms": None,
                    "source_revision": 1,
                    "official_placing": 2,
                },
            ],
            "observed_at_utc": observed_at or "2026-10-02T21:01:00.000Z",
            "deadline_ms": 10000,
        },
    )


def test_actual_numerics_approval_issue_settlement_restart_and_next_round(competition):
    runtime, context, request, root = competition
    forecast_request = {
        key: value
        for key, value in request.items()
        if key not in {"field_id", "upstream_field_revision", "stand_ids", "ceiling", "field_kind"}
    }
    seeds = invoke(runtime, context, "forecast", forecast_request)
    assert seeds["purpose"] == "pre_field_seeding_only" and seeds["issued_mark"] is False
    assert "marks" not in seeds["numeric"]
    receipt = invoke(runtime, context, "field", request)
    assert min(receipt["numeric"]["marks"]) == 3
    assert (
        receipt["numeric"]["forecasts"][0]["predicted_time_ms"]
        < receipt["numeric"]["forecasts"][1]["predicted_time_ms"]
    )
    issue = approve_and_issue(runtime, context, receipt)
    assert issue["issued"] is True
    settlement = settle(runtime, context, receipt, issue)
    assert settlement["same_round_epoch_changed"] is False
    same_round = invoke(
        runtime, context, "field", {**request, "field_id": "field:synthetic-heat-2"}
    )
    assert same_round["numeric"]["epoch_digest"] == receipt["numeric"]["epoch_digest"]
    assert same_round["numeric"]["forecasts"] == receipt["numeric"]["forecasts"]
    second_issue = approve_and_issue(runtime, context, same_round)
    settle(runtime, context, same_round, second_issue)
    closure = invoke(
        runtime,
        context,
        "close_round",
        {
            "schema_version": "strathmark-v3-round-close-request-v1",
            "round_id": request["round_id"],
            "closed_at_utc": "2026-10-02T21:02:00.000Z",
            "deadline_ms": 10000,
        },
    )
    assert closure["status"] == "closed"
    restarted = LinuxCompetitionRuntime(root=root, ml_bundle=runtime.bundle_root)
    assert invoke(restarted, context, "field", request) == receipt
    final = invoke(
        restarted,
        context,
        "field",
        {
            **request,
            "field_id": "field:synthetic-final",
            "round_id": "round:synthetic-final",
            "round_ordinal": 2,
            "predecessor_round_ids": [request["round_id"]],
        },
    )
    assert final["numeric"]["epoch_digest"] != receipt["numeric"]["epoch_digest"]
    assert final["numeric"]["forecasts"] != receipt["numeric"]["forecasts"]
    assert final["numeric"]["weights"] != {"formula": "0.5", "ml": "0.5"}
    assert all(item["capability_adjustments"] for item in final["numeric"]["forecasts"])
    final_issue = approve_and_issue(restarted, context, final)
    settle(restarted, context, final, final_issue)
    invoke(
        restarted,
        context,
        "close_round",
        {
            "schema_version": "strathmark-v3-round-close-request-v1",
            "round_id": "round:synthetic-final",
            "closed_at_utc": "2026-10-02T21:03:00.000Z",
            "deadline_ms": 10000,
        },
    )
    closed = invoke(
        restarted,
        context,
        "close_scope",
        {
            "schema_version": "strathmark-v3-scope-close-request-v1",
            "scope_id": context["scope_id"],
            "closed_at_utc": "2026-10-02T21:04:00.000Z",
            "deadline_ms": 10000,
        },
    )
    assert closed["status"] == "closed"
    backup = restarted.store.backup(root.parent / "independent-recovery")
    assert backup["verified"]
    restored = LinuxCompetitionRuntime(
        root=root.parent / "independent-recovery", ml_bundle=runtime.bundle_root
    )
    assert invoke(restored, context, "field", request) == receipt


def test_issue_before_approval_and_issued_recalculation_fail_without_mutation(competition):
    runtime, context, request, _root = competition
    receipt = invoke(runtime, context, "field", request)
    payload = {
        "schema_version": "strathmark-v3-issue-acknowledgment-request-v1",
        "upstream_issue_id": "issue:synthetic-premature",
        "receipt_bindings": [
            {"receipt_id": receipt["receipt_id"], "receipt_digest": receipt["receipt_digest"]}
        ],
        "issued_at_utc": "2026-10-02T21:00:01.000Z",
        "deadline_ms": 10000,
    }
    before = runtime.store.state()
    with pytest.raises(LinuxLifecycleError, match="prior exact judge approval"):
        invoke(runtime, context, "issue", payload)
    assert runtime.store.state() == before
    approve_and_issue(runtime, context, receipt)
    with pytest.raises(LinuxLifecycleError, match="issued field cannot"):
        invoke(runtime, context, "field", {**request, "upstream_field_revision": 2})
    with pytest.raises(LinuxLifecycleError, match="earlier rounds must close|predecessor round"):
        invoke(
            runtime,
            context,
            "field",
            {
                **request,
                "field_id": "field:synthetic-premature-final",
                "round_id": "round:synthetic-final",
                "round_ordinal": 2,
                "predecessor_round_ids": [request["round_id"]],
            },
        )


def test_complete_issued_field_can_issue_and_settle_without_model_files(competition, tmp_path):
    runtime, context, request, root = competition
    receipt = invoke(runtime, context, "field", request)
    unavailable = LinuxCompetitionRuntime(root=root, ml_bundle=tmp_path / "missing-model")
    issue = approve_and_issue(unavailable, context, receipt)
    assert settle(unavailable, context, receipt, issue)["results_settled"] == 2
    assert invoke(unavailable, context, "field", request) == receipt


def test_correction_retains_issued_receipt_old_results_and_frozen_later_epoch(competition):
    runtime, context, request, _root = competition
    receipt = invoke(runtime, context, "field", request)
    issue = approve_and_issue(runtime, context, receipt)
    original = settle(runtime, context, receipt, issue)
    invoke(
        runtime,
        context,
        "advance",
        {
            "round_ordinal": 2,
            "epoch_group_id": request["epoch_group_id"],
            "closed_at_utc": "2026-10-02T21:02:00.000Z",
        },
    )
    next_request = {
        **request,
        "round_id": "round:synthetic-next",
        "round_ordinal": 2,
        "field_id": "field:synthetic-next",
    }
    later = invoke(runtime, context, "field", next_request)
    before = runtime.store.state()["roots"][context["scope_id"]]
    field = before["fields"][request["field_id"]]
    payload = {
        "schema_version": "strathmark-v3-settlement-request-v1",
        "receipt_id": receipt["receipt_id"],
        "issue_batch_id": issue["issue_batch_id"],
        "observed_at_utc": "2026-10-02T21:03:00.000Z",
        "deadline_ms": 10000,
        "supersedes_settlement_id": original["settlement_id"],
        "reason_code": "official_judge_correction",
        "results": [
            {
                "competitor_id": identifier,
                "status": "completion" if index == 0 else "dnf",
                "raw_time_ms": 24000 if index == 0 else None,
                "penalty_ms": None,
                "source_revision": 2,
                "official_placing": 1 if index == 0 else None,
            }
            for index, identifier in enumerate(receipt["competitor_ids"])
        ],
    }
    corrected = invoke(runtime, context, "correct", payload)
    after = runtime.store.state()["roots"][context["scope_id"]]
    current = after["fields"][request["field_id"]]
    assert current["receipt"] == field["receipt"] and current["issue"] == field["issue"]
    assert current["settlement_history"] == [field["settlement"]]
    assert len(after["reversed_scores"]) == 4
    assert after["phases"] == before["phases"]
    assert invoke(runtime, context, "field", next_request) == later
    assert corrected["source_revision"] == 2 and corrected["frozen_epochs_rewritten"] is False
    assert invoke(runtime, context, "correct", payload) == corrected
    stale = {**payload, "reason_code": "another_reason"}
    with pytest.raises(LinuxLifecycleError, match="latest settlement"):
        invoke(runtime, context, "correct", stale)
    assert runtime.store.state()["roots"][context["scope_id"]] == after


@pytest.mark.parametrize(
    "operation", ["approve", "issue", "settle", "correct", "close_round", "close_scope"]
)
def test_unknown_mutation_schema_rejected_before_any_mutation(competition, operation):
    runtime, context, _request, _root = competition
    before = runtime.store.state()
    with pytest.raises(LinuxLifecycleError, match="schema version"):
        invoke(runtime, context, operation, {"schema_version": "unrecognized-v99"})
    assert runtime.store.state() == before


def test_future_outcomes_rejected_and_judge_ties_retained(competition):
    runtime, context, request, _root = competition
    receipt = invoke(runtime, context, "field", request)
    issue = approve_and_issue(runtime, context, receipt)
    before = runtime.store.state()
    with pytest.raises(LinuxLifecycleError, match="future observed result"):
        settle(runtime, context, receipt, issue, observed_at="2030-01-01T00:00:00.000Z")
    assert runtime.store.state() == before
    marks = receipt["numeric"]["marks"]
    clock = max(marks) * 1000 + 40000
    results = [
        {
            "competitor_id": identifier,
            "status": "completion",
            "raw_time_ms": clock - mark * 1000,
            "penalty_ms": None,
            "source_revision": 1,
            "official_placing": 1,
        }
        for identifier, mark in zip(receipt["competitor_ids"], marks, strict=True)
    ]
    accepted = settle(runtime, context, receipt, issue, results=results)
    assert accepted["legal_finish_order"] is None
    assert accepted["legal_finish_groups"] == [["SYN001", "SYN002"]]
    assert accepted["placing_status"] == "judge_authorized"


def test_events_advance_independently_and_forecast_only_rounds_close(competition):
    runtime, context, request, _root = competition
    forecast = {
        key: value
        for key, value in request.items()
        if key not in {"field_id", "field_kind", "upstream_field_revision", "stand_ids", "ceiling"}
    }
    invoke(runtime, context, "forecast", {**forecast, "round_id": "round:synthetic-seeding"})
    first = invoke(runtime, context, "field", request)
    other_request = {
        **request,
        "epoch_group_id": "round:synthetic-other-event",
        "round_id": "round:synthetic-other-heat",
        "field_id": "field:synthetic-other-heat",
    }
    other = invoke(runtime, context, "field", other_request)
    settle(runtime, context, first, approve_and_issue(runtime, context, first))
    invoke(
        runtime,
        context,
        "advance",
        {
            "round_ordinal": 2,
            "epoch_group_id": request["epoch_group_id"],
            "closed_at_utc": "2026-10-02T21:02:00.000Z",
        },
    )
    final = invoke(
        runtime,
        context,
        "field",
        {
            **request,
            "round_id": "round:synthetic-final",
            "field_id": "field:synthetic-final",
            "round_ordinal": 2,
        },
    )
    root = runtime.store.state()["roots"][context["scope_id"]]
    assert root["rounds"][other_request["round_id"]]["status"] == "open"
    other_again = invoke(
        runtime,
        context,
        "field",
        {**other_request, "field_id": "field:synthetic-other-second-heat"},
    )
    assert other_again["numeric"]["epoch_digest"] == other["numeric"]["epoch_digest"]
    assert other_again["numeric"]["forecasts"] == other["numeric"]["forecasts"]
    for receipt in (final, other, other_again):
        settle(runtime, context, receipt, approve_and_issue(runtime, context, receipt))
    for round_id in (final["round_id"], other["round_id"]):
        invoke(
            runtime,
            context,
            "close_round",
            {
                "schema_version": "strathmark-v3-round-close-request-v1",
                "round_id": round_id,
                "closed_at_utc": "2026-10-02T21:03:00.000Z",
                "deadline_ms": 10000,
            },
        )
    # Distinct final seeding identity is not a missing official settlement.
    invoke(
        runtime,
        context,
        "forecast",
        {**forecast, "round_id": "round:synthetic-final-seeding", "round_ordinal": 2},
    )
    closed = invoke(
        runtime,
        context,
        "close_scope",
        {
            "schema_version": "strathmark-v3-scope-close-request-v1",
            "scope_id": context["scope_id"],
            "closed_at_utc": "2026-10-02T21:04:00.000Z",
            "deadline_ms": 10000,
        },
    )
    assert closed["status"] == "closed"
    assert all(
        value["status"] == "closed"
        for value in runtime.store.state()["roots"][context["scope_id"]]["rounds"].values()
    )
