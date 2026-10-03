"""Complete local Linux competition lifecycle with real Formula and trained ML.

This separate, versioned profile is authorized by the local competition operator.
It does not advertise the Windows CNG production qualification or change V2.
No network service, external provider, or fresh model is needed to issue or settle
an already prepared field. All state transitions use signed exact-retry authority.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import datetime, timezone
from decimal import Decimal
from hashlib import sha256
from pathlib import Path

from strathmark.v3.contracts.canonical import canonical_decimal_string
from strathmark.v3.contracts.evidence import (
    ResultObservation,
    TargetContext,
    require_utc_milliseconds,
)
from strathmark.v3.contracts.forecasts import PositiveTimeDistribution
from strathmark.v3.contracts.identifiers import StableIdentifier, deterministic_identifier
from strathmark.v3.contracts.statuses import OfficialResult, ResultStatus, admit_raw_completion
from strathmark.v3.domain.credibility import compute_predictive_metrics
from strathmark.v3.factory.ml_artifacts import load_ml_bundle
from strathmark.v3.infrastructure.integrity import sign_manifest
from strathmark.v3.linux_forecasts import POLICY, calculate, load_formula_manifest
from strathmark.v3.linux_lifecycle_store import (
    MAX_BYTES,
    LinuxLifecycleError,
    LinuxLifecycleStore,
    digest,
    encode,
)
from strathmark.v3.runtime_identity import implementation_digest

PROTOCOL = "strathmark.v3-linux-competition.v1"
OPERATIONS = (
    "forecast",
    "field",
    "approval_page",
    "approval_detail",
    "approve",
    "issue",
    "settle",
    "correct",
    "close_round",
    "close_scope",
    "advance",
    "lookup",
)
CONTRACT_DIGEST = digest(
    {
        "protocol": PROTOCOL,
        "operations": list(OPERATIONS),
        "forecast_policy": POLICY,
        "authority": "local_competition_operator",
        "same_round_learning": False,
        "approval_is_issue": False,
        "issued_marks_mutable": False,
        "field_kinds": ["handicap", "championship"],
        "placings": "judge_authorized_or_explicitly_unresolved",
        "epoch_groups": "event_local_prior_rounds",
        "mutation_schemas": "exact_version_required",
        "competitor_bindings": "scope_bijective_immutable",
        "counterfactual_marks": "complete_field_optimizer",
    }
)
CONTEXT_FIELDS = {
    "scope_id",
    "selected_engine",
    "mode",
    "contract_identity",
    "source_identity",
    "selected_by_actor_id",
    "locked_at",
}
FORECAST_FIELDS = {
    "workbook",
    "cutoff_at_utc",
    "round_id",
    "round_ordinal",
    "epoch_group_id",
    "predecessor_round_ids",
    "competitor_ids",
    "upstream_competitor_ids",
    "target_context",
}
FIELD_FIELDS = FORECAST_FIELDS | {
    "field_id",
    "upstream_field_revision",
    "stand_ids",
    "ceiling",
    "field_kind",
}

SCHEMAS = {
    "approve": "strathmark-v3-approval-decision-request-v1",
    "issue": "strathmark-v3-issue-acknowledgment-request-v1",
    "settle": "strathmark-v3-settlement-request-v1",
    "correct": "strathmark-v3-settlement-request-v1",
    "close_round": "strathmark-v3-round-close-request-v1",
    "close_scope": "strathmark-v3-scope-close-request-v1",
}


def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def _fields(payload, expected):
    if not isinstance(payload, dict) or set(payload) != expected:
        raise LinuxLifecycleError("operation request differs from its closed Linux contract")


def _positive(value, label):
    if type(value) is not int or value < 1:
        raise LinuxLifecycleError(f"{label} must be a positive integer")


def _identifier(value, namespace):
    identifier = StableIdentifier(value)
    if not str(identifier).startswith(namespace + ":"):
        raise LinuxLifecycleError(f"{namespace} identity is required")
    return identifier


class LinuxCompetitionRuntime:
    def __init__(self, *, root: Path, ml_bundle: Path) -> None:
        if sys.version_info[:2] != (3, 13):
            raise LinuxLifecycleError("the Linux competition runtime requires Python 3.13")
        self.store = LinuxLifecycleStore(root)
        self.bundle_root = Path(ml_bundle).resolve(strict=False)

    def _identity(self, bundle_digest):
        from strathmark import __version__

        return {
            "source_digest": implementation_digest(),
            "package_version": __version__,
            "formula_digest": load_formula_manifest(self.bundle_root).digest,
            "ml_bundle_digest": bundle_digest,
            "installation_identity": self.store.signer.identity.to_dict(),
        }

    def status(self) -> dict:
        try:
            import catboost

            bundle = load_ml_bundle(
                self.bundle_root,
                installed_catboost_version=catboost.__version__,
                installed_python_abi="cp313",
            )
            bundle_digest, numeric_available = bundle.digest, True
        except (ImportError, OSError, ValueError, RuntimeError):
            roots = self.store.state()["roots"]
            if not roots:
                raise LinuxLifecycleError(
                    "verified ML bundle unavailable; no retained competition identity exists"
                ) from None
            bundle_digest = sorted(
                root["runtime_identity"]["ml_bundle_digest"] for root in roots.values()
            )[0]
            numeric_available = False
        identity = self._identity(bundle_digest)
        return {
            **identity,
            "source_identity": digest(identity),
            "protocol": PROTOCOL,
            "contract_digest": CONTRACT_DIGEST,
            "available": True,
            "numeric_available": numeric_available,
            "competition_sources": {
                scope: digest(self._identity(root["runtime_identity"]["ml_bundle_digest"]))
                for scope, root in self.store.state()["roots"].items()
            },
            "purpose": "competition_lifecycle",
            "mode": "local",
            "windows_production_qualified": False,
            "assessors": {
                "formula": "available",
                "ml": "available" if numeric_available else "unavailable",
                "llm_council": "unavailable",
            },
            "operations": list(OPERATIONS),
        }

    def _context(self, context):
        _fields(context, CONTEXT_FIELDS)
        _identifier(context["scope_id"], "tournament")
        require_utc_milliseconds(context["locked_at"])
        if (
            context["selected_engine"] != "v3"
            or context["mode"] != "local"
            or context["contract_identity"] != CONTRACT_DIGEST
        ):
            raise LinuxLifecycleError(
                "request differs from the selected Linux competition authority"
            )
        if (
            not isinstance(context["selected_by_actor_id"], str)
            or not context["selected_by_actor_id"]
        ):
            raise LinuxLifecycleError("competition operator identity is required")
        root = self.store.state()["roots"].get(context["scope_id"])
        # Prepared receipt recovery, approval, issue, and settlement use retained
        # model identity, not a fresh inference dependency. New forecasting still
        # requires the complete verified model bundle below.
        expected = (
            self.status()["source_identity"]
            if root is None
            else digest(self._identity(root["runtime_identity"]["ml_bundle_digest"]))
        )
        if context["source_identity"] != expected:
            raise LinuxLifecycleError(
                "installed code, model, or signing identity differs from this competition"
            )

    def execute(self, operation: str, envelope: dict) -> dict:
        if operation not in OPERATIONS:
            raise LinuxLifecycleError("unsupported Linux competition operation")
        _fields(envelope, {"command_id", "context", "payload"})
        context, payload = envelope["context"], envelope["payload"]
        if operation in SCHEMAS and (
            not isinstance(payload, dict) or payload.get("schema_version") != SCHEMAS[operation]
        ):
            raise LinuxLifecycleError("operation schema version differs from the selected contract")
        self._context(context)
        occurred_at = now()
        request = {"context": context, "payload": payload}
        # Recover committed responses before doing forecasting or re-evaluating
        # an approval snapshot that has legitimately changed since the command.
        recovered = self.store.lookup(envelope["command_id"], request, operation=operation)
        if recovered is not None:
            return recovered
        if operation in {"approval_page", "approval_detail", "lookup"}:
            state = self.store.state()
            root = self._root(state, context)
            if operation == "approval_page":
                _fields(payload, {"offset", "limit"})
                return self._page(root, **payload)
            if operation == "approval_detail":
                _fields(payload, {"snapshot_id", "receipt_id"})
                page = self._page(root, offset=0, limit=10000)
                if page["snapshot_id"] != payload["snapshot_id"]:
                    raise LinuxLifecycleError("approval snapshot is stale")
                field = self._receipt(root, payload["receipt_id"])
                return {"receipt_id": field["receipt"]["receipt_id"], "detail": field["receipt"]}
            _fields(payload, {"receipt_id"})
            return self._receipt(root, payload["receipt_id"])["receipt"]
        if operation in {"forecast", "field"}:
            return self._prepare(operation, envelope, context, payload, occurred_at)

        def transition(state):
            root = self._root(state, context)
            if root["status"] != "open" and operation != "correct":
                raise LinuxLifecycleError("competition is closed")
            return {
                "approve": self._approve,
                "issue": self._issue,
                "settle": self._settle,
                "correct": self._correct,
                "close_round": self._close_round,
                "close_scope": self._close_scope,
                "advance": self._advance,
            }[operation](root, payload, occurred_at)

        return self.store.execute(
            command_id=envelope["command_id"],
            operation=operation,
            scope_id=context["scope_id"],
            request=request,
            occurred_at=occurred_at,
            transition=transition,
        )

    @staticmethod
    def _root(state, context):
        root = state["roots"].get(context["scope_id"])
        if root is None or root["selection"] != context:
            raise LinuxLifecycleError("competition does not match its immutable selection")
        return root

    def _history_snapshot(self, context, payload):
        folder = self.store.root / "history" / sha256(context["scope_id"].encode()).hexdigest()
        folder.mkdir(parents=True, mode=0o700, exist_ok=True)
        path = folder / "history.xlsx"
        metadata_path = folder / "history.json"
        if metadata_path.exists():
            metadata = json.loads(metadata_path.read_bytes())
            if (
                metadata["selection"] != context
                or metadata["cutoff_at_utc"] != payload["cutoff_at_utc"]
                or sha256(path.read_bytes()).hexdigest() != metadata["sha256"]
            ):
                raise LinuxLifecycleError("competition history snapshot differs")
        else:
            if path.exists():
                raise LinuxLifecycleError("uncommitted history snapshot requires recovery")
            source = Path(payload["workbook"]).resolve(strict=True)
            raw = source.read_bytes()
            if len(raw) > 100_000_000:
                raise LinuxLifecycleError("history workbook exceeds capacity")
            metadata = {
                "selection": context,
                "cutoff_at_utc": payload["cutoff_at_utc"],
                "sha256": sha256(raw).hexdigest(),
                "path": str(path.relative_to(self.store.root)),
            }
            # Unique scope writes are exclusive. A partial snapshot is rejected;
            # it is never silently replaced with changing operator history.
            descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o400)
            with os.fdopen(descriptor, "wb") as stream:
                stream.write(raw)
                stream.flush()
                os.fsync(stream.fileno())
            if sha256(source.read_bytes()).hexdigest() != metadata["sha256"]:
                raise LinuxLifecycleError("operator history changed during snapshot creation")
            # Local issued results are admitted from the signed settlement log,
            # never a second time as unauthenticated legacy workbook history.
            from openpyxl import load_workbook

            workbook = load_workbook(path)
            removed = 0
            try:
                for sheet in workbook.worksheets:
                    if sheet.title.lower() != "results":
                        continue
                    headers = [cell.value for cell in sheet[1]]
                    if "V3 Receipt" not in headers:
                        continue
                    column = headers.index("V3 Receipt") + 1
                    known = {
                        field["receipt"]["receipt_id"]
                        for root in self.store.state()["roots"].values()
                        for field in root["fields"].values()
                        if field.get("settlement")
                    }
                    for index in range(sheet.max_row, 1, -1):
                        receipt_id = sheet.cell(index, column).value
                        if receipt_id:
                            if receipt_id not in known:
                                raise LinuxLifecycleError(
                                    "workbook has issued V3 results whose signed authority is missing; restore the original runtime"
                                )
                            sheet.delete_rows(index)
                            removed += 1
                if removed:
                    path.chmod(0o600)
                    workbook.save(path)
                    with path.open("rb+") as stream:
                        os.fsync(stream.fileno())
                    path.chmod(0o400)
            finally:
                workbook.close()
            metadata.update(
                source_workbook_sha256=metadata["sha256"],
                sha256=sha256(path.read_bytes()).hexdigest(),
                signed_rows_excluded_from_legacy_history=removed,
            )
            descriptor = os.open(metadata_path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
            with os.fdopen(descriptor, "wb") as stream:
                stream.write(encode(metadata))
                stream.flush()
                os.fsync(stream.fileno())
        return metadata

    def _freeze(self, envelope, context, payload, occurred_at):
        scope, round_id = context["scope_id"], payload["round_id"]
        state = self.store.state()
        existing = state["roots"].get(scope)
        if existing is not None:
            self._root(state, context)
            history = existing["history"]
            if payload["cutoff_at_utc"] != history["cutoff_at_utc"]:
                raise LinuxLifecycleError("competition historical cutoff cannot change")
        else:
            history = self._history_snapshot(context, payload)
        freeze_request = {
            "context": context,
            "round_id": round_id,
            "round_ordinal": payload["round_ordinal"],
            "epoch_group_id": payload["epoch_group_id"],
            "predecessor_round_ids": payload["predecessor_round_ids"],
            "history_sha256": history["sha256"],
            "cutoff_at_utc": history["cutoff_at_utc"],
        }

        binding = self.status()
        identity = {
            key: binding[key]
            for key in (
                "source_digest",
                "package_version",
                "formula_digest",
                "ml_bundle_digest",
                "installation_identity",
            )
        }

        def transition(state):
            root = state["roots"].setdefault(
                scope,
                {
                    "selection": context,
                    "runtime_identity": identity,
                    "history": history,
                    "status": "open",
                    "rounds": {},
                    "phases": {},
                    "fields": {},
                    "forecasts": {},
                    "decisions": {},
                    "scores": [],
                },
            )
            if root["selection"] != context or root["status"] != "open":
                raise LinuxLifecycleError("competition selection changed or is closed")
            if round_id in root["rounds"]:
                raise LinuxLifecycleError("round already exists with another freeze identity")
            ordinal = payload["round_ordinal"]
            group = payload["epoch_group_id"]
            group_rounds = {
                key: value
                for key, value in root["rounds"].items()
                if value["epoch_group_id"] == group
            }
            if group_rounds and ordinal < max(value["ordinal"] for value in group_rounds.values()):
                raise LinuxLifecycleError(
                    "cannot create an earlier-round field after a later epoch has started"
                )
            predecessors = payload["predecessor_round_ids"]
            if not predecessors and payload["round_ordinal"] > 1:
                predecessors = [
                    key
                    for key, value in group_rounds.items()
                    if value["ordinal"] < payload["round_ordinal"]
                ]
                if not predecessors:
                    raise LinuxLifecycleError("later round has no completed predecessor")
            if any(
                key not in group_rounds or group_rounds[key]["status"] != "closed"
                for key in predecessors
            ):
                raise LinuxLifecycleError("predecessor round is not closed and settled")
            phase_key = group + ":" + str(ordinal)
            phase = root["phases"].get(phase_key)
            if phase is None:
                if any(
                    value["ordinal"] < ordinal and value["status"] != "closed"
                    for value in group_rounds.values()
                ):
                    raise LinuxLifecycleError(
                        "all fields in earlier rounds must close before the next epoch"
                    )
                live = []
                for other_scope, other_root in sorted(state["roots"].items()):
                    if other_scope != scope and (
                        other_root["status"] != "closed"
                        or other_root["closed_at_utc"] >= history["cutoff_at_utc"]
                    ):
                        continue
                    for field in other_root["fields"].values():
                        source_round = other_root["rounds"][field["receipt"]["round_id"]]
                        if (
                            not field.get("settlement")
                            or source_round["status"] != "closed"
                            or (other_scope == scope and source_round["ordinal"] >= ordinal)
                        ):
                            continue
                        live.extend(
                            item
                            for item in field["settlement"]["live_results"]
                            if item["observation"]["occurred_at_utc"] <= occurred_at
                            and (
                                other_scope == scope
                                or item["observation"]["occurred_at_utc"] < history["cutoff_at_utc"]
                            )
                        )
                live.sort(
                    key=lambda item: (
                        item["observation"]["occurred_at_utc"],
                        item["observation"]["evidence_id"],
                    )
                )
                phase = {
                    "live_results": live,
                    "weights": self._weights(root, ordinal),
                    "frozen_at_utc": occurred_at,
                    "epoch_round_id": str(
                        deterministic_identifier(
                            "round", {"scope": scope, "epoch_group_id": group, "ordinal": ordinal}
                        )
                    ),
                }
                root["phases"][phase_key] = phase
            root["rounds"][round_id] = {
                "status": "open",
                "ordinal": payload["round_ordinal"],
                "epoch_group_id": group,
                "predecessors": predecessors,
                "history_path": history["path"],
                "history_sha256": history["sha256"],
                "cutoff_at_utc": history["cutoff_at_utc"],
                **phase,
            }
            root["rounds"][round_id]["snapshot_digest"] = digest(root["rounds"][round_id])
            return {
                "round_id": round_id,
                "epoch_snapshot_digest": root["rounds"][round_id]["snapshot_digest"],
            }

        key = "freeze:" + digest({"scope": scope, "round": round_id})
        recovered = self.store.lookup(key, freeze_request)
        if recovered is None:
            self.store.execute(
                command_id=key,
                operation="freeze_round",
                scope_id=scope,
                request=freeze_request,
                occurred_at=occurred_at,
                transition=transition,
            )
        return self._root(self.store.state(), context)["rounds"][round_id]

    @staticmethod
    def _weights(root, ordinal):
        # A disclosed Linux policy, not a replacement for the Windows factory:
        # eight equal-prior opportunities regularize earned normalized CRPS.
        values = {}
        for assessor in ("formula", "ml"):
            scores = [
                Decimal(item["normalized_crps"])
                for item in root["scores"]
                if item["assessor"] == assessor
                and root["rounds"][item["round_id"]]["status"] == "closed"
                and root["rounds"][item["round_id"]]["ordinal"] < ordinal
            ]
            loss = (sum(scores) + Decimal("0.8")) / (len(scores) + 8)
            values[assessor] = Decimal(1) / max(loss, Decimal("0.001"))
        formula = values["formula"] / sum(values.values())
        formula = min(Decimal("0.9"), max(Decimal("0.1"), formula))
        return {
            "formula": canonical_decimal_string(formula),
            "ml": canonical_decimal_string(1 - formula),
        }

    @staticmethod
    def _competitor_bindings(root, payload):
        if root is None:
            return
        upstream_to_local, local_to_upstream = {}, {}
        receipts = list(root["forecasts"].values()) + [
            field["receipt"] for field in root["fields"].values()
        ]
        receipts += [field["receipt"] for field in root.get("superseded_fields", {}).values()]
        for receipt in receipts:
            for upstream, local in zip(
                receipt["competitor_ids"], receipt["local_competitor_ids"], strict=True
            ):
                upstream_to_local[upstream], local_to_upstream[local] = local, upstream
        for upstream, local in zip(
            payload["upstream_competitor_ids"], payload["competitor_ids"], strict=True
        ):
            if (
                upstream_to_local.get(upstream, local) != local
                or local_to_upstream.get(local, upstream) != upstream
            ):
                raise LinuxLifecycleError(
                    "competitor identity binding cannot change within a competition"
                )

    def _prepare(self, operation, envelope, context, payload, occurred_at):
        _fields(payload, FORECAST_FIELDS if operation == "forecast" else FIELD_FIELDS)
        _identifier(payload["round_id"], "round")
        _identifier(payload["epoch_group_id"], "round")
        _positive(payload["round_ordinal"], "round ordinal")
        require_utc_milliseconds(payload["cutoff_at_utc"])
        if not isinstance(payload["predecessor_round_ids"], list) or len(
            set(payload["predecessor_round_ids"])
        ) != len(payload["predecessor_round_ids"]):
            raise LinuxLifecycleError("round predecessors must be distinct")
        for predecessor in payload["predecessor_round_ids"]:
            _identifier(predecessor, "round")
        TargetContext.from_dict(payload["target_context"])
        upstream = payload["upstream_competitor_ids"]
        if (
            not isinstance(upstream, list)
            or len(upstream) != len(payload["competitor_ids"])
            or len(set(upstream)) != len(upstream)
        ):
            raise LinuxLifecycleError("upstream competitor bindings differ from the exact roster")
        for identifier in upstream:
            _identifier(identifier, "competitor")
        if operation == "field":
            if payload["field_kind"] not in {"handicap", "championship"}:
                raise LinuxLifecycleError("explicit handicap or championship field kind required")
            _identifier(payload["field_id"], "field")
            _positive(payload["upstream_field_revision"], "field revision")
            if type(payload["ceiling"]) is not int or not 3 <= payload["ceiling"] <= 180:
                raise LinuxLifecycleError("mark ceiling is invalid")
            if (
                not isinstance(payload["stand_ids"], list)
                or len(payload["stand_ids"]) != len(payload["competitor_ids"])
                or len(set(payload["stand_ids"])) != len(payload["stand_ids"])
            ):
                raise LinuxLifecycleError("field requires distinct stands for its exact roster")
            for stand in payload["stand_ids"]:
                _identifier(stand, "stand")
        self._competitor_bindings(self.store.state()["roots"].get(context["scope_id"]), payload)
        frozen = self._freeze(envelope, context, payload, occurred_at)
        if frozen["status"] != "open":
            raise LinuxLifecycleError("round is closed")
        before = self.store.state()
        root = self._root(before, context)
        previous = root["fields"].get(payload.get("field_id"))
        if previous and (previous.get("issue") or previous.get("settlement")):
            raise LinuxLifecycleError("issued field cannot be recalculated or revised")
        if (
            previous
            and payload["upstream_field_revision"] <= previous["receipt"]["upstream_field_revision"]
        ):
            raise LinuxLifecycleError("field revision must increase for changed work")
        readiness = self.status()
        if (
            readiness["numeric_available"] is not True
            or readiness["source_identity"] != context["source_identity"]
        ):
            raise LinuxLifecycleError("model identity changed before new forecasting")
        numeric_snapshot = {**frozen, "history_path": str(self.store.root / frozen["history_path"])}
        numeric = calculate(
            bundle_root=self.bundle_root,
            signer=self.store.signer,
            trust=self.store.trust,
            scope_id=context["scope_id"],
            round_id=frozen["epoch_round_id"],
            round_snapshot=numeric_snapshot,
            competitor_ids=payload["competitor_ids"],
            target_context=payload["target_context"],
            field_id=payload.get("field_id")
            if payload.get("field_kind") != "championship"
            else None,
            ceiling=payload.get("ceiling", 180),
        )
        if operation == "field" and payload["field_kind"] == "championship":
            numeric.update(
                marks=[3] * len(payload["competitor_ids"]),
                mark_policy="championship_fixed_mark_3",
                counterfactual_marks={
                    key: [3] * len(payload["competitor_ids"]) for key in ("formula", "ml")
                },
                maximum_mark_disagreement=0,
                classification="amber",
            )
        receipt = {
            "protocol": PROTOCOL,
            "scope_id": context["scope_id"],
            "round_id": payload["round_id"],
            "source_identity": context["source_identity"],
            "selection": context,
            "epoch_snapshot_digest": frozen["snapshot_digest"],
            "purpose": "pre_field_seeding_only" if operation == "forecast" else "complete_field",
            "issued_mark": False,
            "numeric": numeric,
            "competitor_ids": upstream,
            "local_competitor_ids": payload["competitor_ids"],
            "target_context": payload["target_context"],
        }
        if operation == "field":
            receipt.update(
                field_kind=payload["field_kind"],
                field_id=payload["field_id"],
                upstream_field_revision=payload["upstream_field_revision"],
                stand_ids=payload["stand_ids"],
                ceiling=payload["ceiling"],
            )
        receipt_id = str(deterministic_identifier("receipt", receipt))
        receipt["receipt_id"] = receipt_id
        receipt["receipt_digest"] = digest(receipt)

        def transition(state):
            current = self._root(state, context)
            self._competitor_bindings(current, payload)
            round_state = current["rounds"][payload["round_id"]]
            if (
                round_state["status"] != "open"
                or round_state["snapshot_digest"] != frozen["snapshot_digest"]
            ):
                raise LinuxLifecycleError("round authority changed during forecasting")
            if operation == "forecast":
                current["forecasts"][receipt_id] = receipt
            else:
                latest = current["fields"].get(payload["field_id"])
                if latest != previous:
                    raise LinuxLifecycleError("field authority changed during forecasting")
                if previous is not None:
                    current.setdefault("superseded_fields", {})[
                        previous["receipt"]["receipt_id"]
                    ] = previous
                current["fields"][payload["field_id"]] = {
                    "receipt": receipt,
                    "decision": None,
                    "issue": None,
                    "settlement": None,
                }
            return receipt

        return self.store.execute(
            command_id=envelope["command_id"],
            operation=operation,
            scope_id=context["scope_id"],
            request={"context": context, "payload": payload},
            occurred_at=occurred_at,
            transition=transition,
        )

    @staticmethod
    def _receipt(root, receipt_id):
        matches = [
            field
            for field in root["fields"].values()
            if field["receipt"]["receipt_id"] == receipt_id
        ]
        if len(matches) != 1:
            raise LinuxLifecycleError("receipt is absent, superseded, or outside this scope")
        return matches[0]

    def _page(self, root, *, offset, limit):
        if (
            type(offset) is not int
            or offset < 0
            or type(limit) is not int
            or not 1 <= limit <= 10000
        ):
            raise LinuxLifecycleError("approval page capacity is invalid")
        rows = []
        for order, (field_id, field) in enumerate(sorted(root["fields"].items()), 1):
            receipt = field["receipt"]
            classification = receipt["numeric"]["classification"]
            decision = field["decision"]
            row = {
                "field_id": field_id,
                "receipt_id": receipt["receipt_id"],
                "receipt_content_digest": receipt["receipt_digest"],
                "receipt_revision": 1,
                "upstream_field_revision": receipt["upstream_field_revision"],
                "call_order": order,
                "classification": classification,
                "lane": "red" if classification == "red" else "degraded",
                "decision_state": "undecided" if decision is None else decision["action"],
                "ordinary_batch_eligible": False,
                "degraded_batch_eligible": classification != "red",
                "causal_rule_codes": ["llm_council_unavailable", "formula_ml_mark_disagreement"]
                if classification == "red"
                else ["llm_council_unavailable"],
                "affected_competitors": receipt["competitor_ids"],
            }
            row["row_digest"] = digest(row)
            rows.append(row)
        snapshot_id = "snapshot:" + digest(rows)
        return {
            "snapshot_id": snapshot_id,
            "rows": rows[offset : offset + limit],
            "total": len(rows),
        }

    def _approve(self, root, payload, occurred_at):
        expected = {
            "schema_version",
            "tournament_id",
            "snapshot_id",
            "action",
            "selected",
            "excluded",
            "actor_metadata",
            "reason_code",
            "superseded_receipt_id",
            "decided_at_utc",
            "deadline_ms",
        }
        _fields(payload, expected)
        if payload["tournament_id"] != root["selection"]["scope_id"]:
            raise LinuxLifecycleError("approval competition differs")
        require_utc_milliseconds(payload["decided_at_utc"])
        page = self._page(root, offset=0, limit=10000)
        if payload["snapshot_id"] != page["snapshot_id"]:
            raise LinuxLifecycleError("approval snapshot is stale")
        actions = {
            "ordinary_batch_accept",
            "degraded_batch_accept",
            "individual_accept",
            "exclude",
            "defer",
        }
        action = payload["action"]
        if (
            action not in actions
            or not isinstance(payload["reason_code"], str)
            or not payload["reason_code"]
        ):
            raise LinuxLifecycleError("deliberate approval action and reason are required")
        if (
            not isinstance(payload["selected"], list)
            or not payload["selected"]
            or not isinstance(payload["excluded"], list)
        ):
            raise LinuxLifecycleError("approval must bind explicit selected and excluded receipts")
        if action in {"individual_accept", "exclude", "defer"} and len(payload["selected"]) != 1:
            raise LinuxLifecycleError("individual review must select exactly one receipt")
        by_id = {row["receipt_id"]: row for row in page["rows"]}
        seen = set()
        for selected in (*payload["selected"], *payload["excluded"]):
            row = by_id.get(selected.get("receipt_id"))
            if row is None or row["receipt_id"] in seen:
                raise LinuxLifecycleError("approval has duplicate or stale receipt binding")
            expected_binding = {
                "field_id": row["field_id"],
                "receipt_id": row["receipt_id"],
                "receipt_digest": row["receipt_content_digest"],
                "receipt_revision": row["receipt_revision"],
                "upstream_field_revision": row["upstream_field_revision"],
                "row_digest": row["row_digest"],
                "call_order": row["call_order"],
            }
            if selected != expected_binding or row["decision_state"] != "undecided":
                raise LinuxLifecycleError("approval binding differs from its undecided exact row")
            if (
                selected in payload["selected"]
                and action == "ordinary_batch_accept"
                and not row["ordinary_batch_eligible"]
            ):
                raise LinuxLifecycleError(
                    "this profile requires explicit degraded or individual review"
                )
            if (
                selected in payload["selected"]
                and action == "degraded_batch_accept"
                and not row["degraded_batch_eligible"]
            ):
                raise LinuxLifecycleError("red disagreement cannot enter degraded batching")
            seen.add(row["receipt_id"])
        decision_id = "decision:" + digest(payload)
        for binding in payload["selected"]:
            self._receipt(root, binding["receipt_id"])["decision"] = {
                "decision_id": decision_id,
                "action": action,
                "decided_at_utc": payload["decided_at_utc"],
                "reason_code": payload["reason_code"],
            }
        for binding in payload["excluded"]:
            self._receipt(root, binding["receipt_id"])["decision"] = {
                "decision_id": decision_id,
                "action": "exclude",
                "decided_at_utc": payload["decided_at_utc"],
                "reason_code": payload["reason_code"],
            }
        root["decisions"][decision_id] = payload
        return {
            "decision_id": decision_id,
            "action": action,
            "approved_receipt_ids": [item["receipt_id"] for item in payload["selected"]]
            if action.endswith("accept")
            else [],
            "issued": False,
        }

    def _issue(self, root, payload, occurred_at):
        _fields(
            payload,
            {
                "schema_version",
                "upstream_issue_id",
                "receipt_bindings",
                "issued_at_utc",
                "deadline_ms",
            },
        )
        require_utc_milliseconds(payload["issued_at_utc"])
        bindings = payload["receipt_bindings"]
        if (
            not isinstance(bindings, list)
            or not bindings
            or len({item["receipt_id"] for item in bindings}) != len(bindings)
        ):
            raise LinuxLifecycleError("issue requires distinct exact approved receipts")
        fields = []
        for binding in bindings:
            _fields(binding, {"receipt_id", "receipt_digest"})
            field = self._receipt(root, binding["receipt_id"])
            if (
                field["receipt"]["receipt_digest"] != binding["receipt_digest"]
                or field["decision"] is None
                or field["decision"]["action"]
                not in {"ordinary_batch_accept", "degraded_batch_accept", "individual_accept"}
            ):
                raise LinuxLifecycleError("issue requires prior exact judge approval")
            if (
                field["issue"] is not None
                or root["rounds"][field["receipt"]["round_id"]]["status"] != "open"
            ):
                raise LinuxLifecycleError("issue receipt is already issued or round is closed")
            fields.append(field)
        batch_id = "issue_batch:" + digest(payload)
        for field in fields:
            field["issue"] = {
                "issue_batch_id": batch_id,
                "upstream_issue_id": payload["upstream_issue_id"],
                "issued_at_utc": payload["issued_at_utc"],
            }
        return {
            "issue_batch_id": batch_id,
            "receipt_ids": [item["receipt_id"] for item in bindings],
            "issued": True,
            "issued_field_marks": {
                field["receipt"]["receipt_id"]: dict(
                    zip(
                        field["receipt"]["competitor_ids"],
                        field["receipt"]["numeric"]["marks"],
                        strict=True,
                    )
                )
                for field in fields
            },
        }

    def _settle(self, root, payload, occurred_at, *, revision=1):
        _fields(
            payload,
            {
                "schema_version",
                "issue_batch_id",
                "receipt_id",
                "results",
                "observed_at_utc",
                "deadline_ms",
            },
        )
        require_utc_milliseconds(payload["observed_at_utc"])
        if payload["observed_at_utc"] > occurred_at:
            raise LinuxLifecycleError(
                "future observed result cannot enter the causal evidence history"
            )
        field = self._receipt(root, payload["receipt_id"])
        receipt = field["receipt"]
        if field["issue"] is None or payload["issue_batch_id"] != field["issue"]["issue_batch_id"]:
            raise LinuxLifecycleError(
                "settlement requires the exact issued receipt and acknowledgment"
            )
        if payload["observed_at_utc"] < field["issue"]["issued_at_utc"]:
            raise LinuxLifecycleError("result cannot precede its issued sheet")
        if field["settlement"] is not None:
            raise LinuxLifecycleError(
                "settled results are immutable; a correction requires its own revision workflow"
            )
        results = payload["results"]
        if (
            not isinstance(results, list)
            or len(results) != len(receipt["competitor_ids"])
            or {item["competitor_id"] for item in results} != set(receipt["competitor_ids"])
        ):
            raise LinuxLifecycleError(
                "settlement requires one explicit outcome for every issued competitor"
            )
        parsed, placements = {}, {}
        for item in results:
            _fields(
                item,
                {
                    "competitor_id",
                    "status",
                    "raw_time_ms",
                    "penalty_ms",
                    "source_revision",
                    "official_placing",
                },
            )
            _positive(item["source_revision"], "result revision")
            if item["source_revision"] != revision:
                raise LinuxLifecycleError(
                    "result revision must be the exact next official revision"
                )
            official = OfficialResult(
                ResultStatus(item["status"]),
                item["raw_time_ms"],
                item["penalty_ms"],
                item["source_revision"],
                None if revision == 1 else revision - 1,
            )
            local_id = dict(
                zip(receipt["competitor_ids"], receipt["local_competitor_ids"], strict=True)
            )[item["competitor_id"]]
            placing = item["official_placing"]
            if placing is not None:
                _positive(placing, "official placing")
                if official.status not in {ResultStatus.COMPLETION, ResultStatus.PENALTY}:
                    raise LinuxLifecycleError(
                        "nonfinishes and void results cannot carry legal placings"
                    )
            parsed[local_id], placements[local_id] = official, placing
        marks = dict(zip(receipt["local_competitor_ids"], receipt["numeric"]["marks"], strict=True))
        finishes = sorted(
            (result.raw_time_ms + marks[key] * 1000 + (result.penalty_ms or 0), key)
            for key, result in parsed.items()
            if result.status in {ResultStatus.COMPLETION, ResultStatus.PENALTY}
        )
        legal_groups = [
            [key for key in sorted(parsed) if placements[key] == position]
            for position in sorted({place for place in placements.values() if place is not None})
        ]
        unresolved = any(
            result.status in {ResultStatus.COMPLETION, ResultStatus.PENALTY}
            and placements[key] is None
            for key, result in parsed.items()
        )
        legal_order = (
            [group[0] for group in legal_groups]
            if not unresolved and all(len(group) == 1 for group in legal_groups)
            else None
        )
        front = finishes[0][0] if finishes else None
        live_results = []
        issued = {
            "field_id": receipt["field_id"],
            "upstream_field_revision": receipt["upstream_field_revision"],
            "competitor_ids": [item["competitor_id"] for item in receipt["numeric"]["forecasts"]],
            "receipt_id": receipt["receipt_id"],
            "scope_id": receipt["scope_id"],
            "round_id": receipt["round_id"],
            "target_context": receipt["target_context"],
            "issued_marks": [
                [item["competitor_id"], marks[item["local_competitor_id"]]]
                for item in receipt["numeric"]["forecasts"]
            ],
        }
        for index, forecast in enumerate(receipt["numeric"]["forecasts"], 1):
            local_id = forecast["local_competitor_id"]
            result = parsed[local_id]
            completion = admit_raw_completion(result)
            completion_clock = (
                result.raw_time_ms + marks[local_id] * 1000 if completion is not None else None
            )
            observation = ResultObservation(
                deterministic_identifier(
                    "evidence",
                    {
                        "receipt": receipt["receipt_id"],
                        "competitor": forecast["competitor_id"],
                        "revision": result.revision,
                    },
                ),
                StableIdentifier(forecast["competitor_id"]),
                StableIdentifier(receipt["scope_id"]),
                StableIdentifier(receipt["round_id"]),
                StableIdentifier(receipt["field_id"]),
                TargetContext.from_dict(receipt["target_context"]),
                index,
                payload["observed_at_utc"],
                marks[local_id],
                completion_clock,
                placements.get(local_id),
                completion_clock - front
                if completion_clock is not None and front is not None
                else None,
                result,
                digest(
                    {
                        "payload": payload,
                        "competitor_id": local_id,
                        "issued_receipt_digest": receipt["receipt_digest"],
                    }
                ),
            )
            live_results.append(
                {
                    "observation": observation.to_dict(),
                    "issued_field": issued,
                    "upstream_field_revision": receipt["upstream_field_revision"],
                    "authority_digest": digest(field["issue"]),
                    "prior_median_ms": forecast["predicted_time_ms"],
                }
            )
            if completion is not None:
                for assessor in ("formula", "ml"):
                    distribution = PositiveTimeDistribution.from_dict(
                        forecast[assessor]["forecast"]["distribution"]
                    )
                    metrics = compute_predictive_metrics(
                        distribution,
                        actual_time_ms=completion.raw_time_ms,
                        robust_context_scale_ms=max(1, forecast["predicted_time_ms"]),
                    )
                    root["scores"].append(
                        {
                            "assessor": assessor,
                            "round_id": receipt["round_id"],
                            "receipt_id": receipt["receipt_id"],
                            "competitor_id": local_id,
                            **metrics.to_dict(),
                        }
                    )
        settlement_id = "settlement:" + digest(payload)
        field["settlement"] = {
            "settlement_id": settlement_id,
            "results": results,
            "live_results": live_results,
            "observed_at_utc": payload["observed_at_utc"],
            "legal_finish_order": legal_order,
            "legal_finish_groups": legal_groups,
            "placing_status": "unresolved" if unresolved else "judge_authorized",
        }
        return {
            "settlement_id": settlement_id,
            "receipt_id": receipt["receipt_id"],
            "results_settled": len(results),
            "legal_finish_order": legal_order,
            "legal_finish_groups": legal_groups,
            "placing_status": "unresolved" if unresolved else "judge_authorized",
            "same_round_epoch_changed": False,
        }

    def _correct(self, root, payload, occurred_at):
        expected = {
            "schema_version",
            "issue_batch_id",
            "receipt_id",
            "results",
            "observed_at_utc",
            "deadline_ms",
            "supersedes_settlement_id",
            "reason_code",
        }
        _fields(payload, expected)
        field = self._receipt(root, payload["receipt_id"])
        previous = field["settlement"]
        if previous is None or payload["supersedes_settlement_id"] != previous["settlement_id"]:
            raise LinuxLifecycleError("correction must supersede the exact latest settlement")
        if not isinstance(payload["reason_code"], str) or not payload["reason_code"].strip():
            raise LinuxLifecycleError("official correction requires a deliberate reason")
        if payload["observed_at_utc"] < previous["observed_at_utc"]:
            raise LinuxLifecycleError("correction cannot precede its original observation")
        revision = max(item["source_revision"] for item in previous["results"]) + 1
        field.setdefault("settlement_history", []).append(previous)
        retained = []
        for score in root["scores"]:
            if score["receipt_id"] == payload["receipt_id"]:
                root.setdefault("reversed_scores", []).append(
                    {**score, "reversed_by_correction": digest(payload)}
                )
            else:
                retained.append(score)
        root["scores"] = retained
        field["settlement"] = None
        request = {
            key: value
            for key, value in payload.items()
            if key not in {"supersedes_settlement_id", "reason_code"}
        }
        response = self._settle(root, request, occurred_at, revision=revision)
        field["settlement"]["correction_reason"] = payload["reason_code"]
        field["settlement"]["supersedes_settlement_id"] = previous["settlement_id"]
        return {
            **response,
            "source_revision": revision,
            "supersedes_settlement_id": previous["settlement_id"],
            "frozen_epochs_rewritten": False,
            "issued_marks_rewritten": False,
        }

    def _close_round(self, root, payload, occurred_at):
        _fields(payload, {"schema_version", "round_id", "closed_at_utc", "deadline_ms"})
        round_state = root["rounds"].get(payload["round_id"])
        if round_state is None:
            raise LinuxLifecycleError("round is absent")
        if round_state["status"] == "closed":
            return {
                "round_id": payload["round_id"],
                "closure_id": round_state["closure_id"],
                "status": "closed",
                "epoch_snapshot_digest": round_state["snapshot_digest"],
            }
        fields = [
            field
            for field in root["fields"].values()
            if field["receipt"]["round_id"] == payload["round_id"]
        ]
        has_forecast = any(
            item["round_id"] == payload["round_id"] for item in root["forecasts"].values()
        )
        if (not fields and not has_forecast) or any(
            field["settlement"] is None
            and (field["decision"] is None or field["decision"]["action"] != "exclude")
            for field in fields
        ):
            raise LinuxLifecycleError(
                "round closure requires all fields settled or deliberately excluded"
            )
        require_utc_milliseconds(payload["closed_at_utc"])
        closure_id = "round_closure:" + digest(payload)
        round_state.update(
            status="closed", closure_id=closure_id, closed_at_utc=payload["closed_at_utc"]
        )
        return {
            "round_id": payload["round_id"],
            "closure_id": closure_id,
            "status": "closed",
            "epoch_snapshot_digest": round_state["snapshot_digest"],
        }

    def _advance(self, root, payload, occurred_at):
        _fields(payload, {"round_ordinal", "epoch_group_id", "closed_at_utc"})
        _identifier(payload["epoch_group_id"], "round")
        _positive(payload["round_ordinal"], "round ordinal")
        require_utc_milliseconds(payload["closed_at_utc"])
        closures = []
        for round_id, round_state in sorted(root["rounds"].items()):
            if (
                round_state["epoch_group_id"] == payload["epoch_group_id"]
                and round_state["ordinal"] < payload["round_ordinal"]
            ):
                closures.append(
                    self._close_round(
                        root,
                        {
                            "schema_version": "strathmark-v3-round-close-request-v1",
                            "round_id": round_id,
                            "closed_at_utc": payload["closed_at_utc"],
                            "deadline_ms": 10000,
                        },
                        occurred_at,
                    )
                )
        if not closures:
            raise LinuxLifecycleError("next round has no completed predecessor")
        return {
            "round_ordinal": payload["round_ordinal"],
            "epoch_group_id": payload["epoch_group_id"],
            "closures": closures,
        }

    def _close_scope(self, root, payload, occurred_at):
        _fields(payload, {"schema_version", "scope_id", "closed_at_utc", "deadline_ms"})
        require_utc_milliseconds(payload["closed_at_utc"])
        # Seeding is mark-free evidence, not an unissued competition field.
        # The menu gives seeding its own round identity; close those retained
        # forecast-only rounds without inventing issue or result authority.
        for round_id, round_state in root["rounds"].items():
            if round_state["status"] != "closed" and not any(
                field["receipt"]["round_id"] == round_id for field in root["fields"].values()
            ):
                self._close_round(
                    root,
                    {
                        "schema_version": SCHEMAS["close_round"],
                        "round_id": round_id,
                        "closed_at_utc": payload["closed_at_utc"],
                        "deadline_ms": 10000,
                    },
                    occurred_at,
                )
        if (
            payload["scope_id"] != root["selection"]["scope_id"]
            or not root["rounds"]
            or any(item["status"] != "closed" for item in root["rounds"].values())
        ):
            raise LinuxLifecycleError("competition closure requires every round to be closed")
        require_utc_milliseconds(payload["closed_at_utc"])
        root.update(status="closed", closed_at_utc=payload["closed_at_utc"])
        return {"scope_id": payload["scope_id"], "status": "closed"}


def main():
    parser = argparse.ArgumentParser(description="Complete local Linux V3 competition runtime")
    parser.add_argument("operation", choices=("init", "status", "state", "backup", *OPERATIONS))
    parser.add_argument("--runtime-root", type=Path, required=True)
    parser.add_argument("--ml-bundle", type=Path, required=True)
    parser.add_argument(
        "--backup-dir",
        type=Path,
        help="Existing independent recovery directory; accepted mutations are archived and read back before acknowledgment",
    )
    args = parser.parse_args()
    try:
        if args.operation == "init":
            LinuxLifecycleStore.initialize(args.runtime_root)
        runtime = LinuxCompetitionRuntime(root=args.runtime_root, ml_bundle=args.ml_bundle)
        if args.operation in {"init", "status"}:
            response = runtime.status()
        elif args.operation == "backup":
            if args.backup_dir is None:
                raise LinuxLifecycleError("backup requires --backup-dir")
            response = {"archive": str(runtime.store.archive_backup(args.backup_dir))}
        elif args.operation == "state":
            response = runtime.store.state()
        else:
            raw = sys.stdin.buffer.read(MAX_BYTES + 1)
            if len(raw) > MAX_BYTES:
                raise LinuxLifecycleError("Linux competition request exceeds capacity")
            envelope = json.loads(raw)
            value = runtime.execute(args.operation, envelope)
            identity = {
                "protocol": PROTOCOL,
                "operation": args.operation,
                "request_digest": digest(envelope),
                "response_digest": digest(value),
            }
            response = {
                **identity,
                "response": value,
                "manifest": sign_manifest(
                    "linux_competition_response",
                    identity,
                    signer=runtime.store.signer,
                    created_at=now(),
                ).to_dict(),
            }
        if args.backup_dir is not None and args.operation not in {
            "backup",
            "status",
            "state",
            "approval_page",
            "approval_detail",
            "lookup",
        }:
            runtime.store.archive_backup(args.backup_dir)
        sys.stdout.buffer.write(encode(response) + b"\n")
    except Exception as error:
        print(f"V3 Linux competition failed: {type(error).__name__}: {error}", file=sys.stderr)
        raise SystemExit(2) from None


if __name__ == "__main__":
    main()
