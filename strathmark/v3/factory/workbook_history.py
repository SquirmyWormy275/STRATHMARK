"""Read an explicit legacy workbook snapshot without touching operator storage.

Legacy rows lack authenticated issued receipts. They remain historical candidate
evidence, carry their exact workbook/row digest, and never become live issue facts.
Undated rows cannot satisfy a causal cutoff and are excluded.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import ROUND_HALF_EVEN, Decimal
from hashlib import sha256
from pathlib import Path

from strathmark.v3.contracts.canonical import canonical_decimal_string, canonical_digest
from strathmark.v3.contracts.evidence import ContextProperty, ResultObservation, TargetContext
from strathmark.v3.contracts.identifiers import deterministic_identifier
from strathmark.v3.contracts.statuses import OfficialResult, ResultStatus

EVENT_CODES = {"UH": "underhand", "SB": "standing_block"}


@dataclass(frozen=True, slots=True)
class WorkbookHistory:
    source_sha256: str
    observations: tuple[ResultObservation, ...]
    competitor_ids: tuple[tuple[str, str], ...]
    excluded: tuple[tuple[str, int], ...]


def load_workbook_history(path: Path | str, *, cutoff_at_utc: str) -> WorkbookHistory:
    """Load dated completions strictly before the requested exclusive cutoff."""
    from openpyxl import load_workbook

    path = Path(path).resolve(strict=True)
    before = path.read_bytes()
    digest = sha256(before).hexdigest()
    cutoff = datetime.fromisoformat(cutoff_at_utc.replace("Z", "+00:00"))
    if cutoff.tzinfo is None:
        raise ValueError("historical cutoff must include a timezone")
    excluded: Counter[str] = Counter()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {sheet.title.lower(): sheet for sheet in workbook}
        if not {"competitor", "wood", "results"} <= set(sheets):
            raise ValueError("history workbook requires Competitor, Wood, and Results")

        def records(sheet):
            rows = iter(sheet.values)
            columns = next(rows)
            return [(number, dict(zip(columns, row))) for number, row in enumerate(rows, 2)]

        competitors = {}
        for _, row in records(sheets["competitor"]):
            local_id = str(row.get("CompetitorID") or "").strip()
            if not local_id:
                continue
            if local_id in competitors:
                raise ValueError("competitor sheet repeats a stable ID")
            competitors[local_id] = deterministic_identifier("competitor", {"local_id": local_id})
        densities = {}
        for _, row in records(sheets["wood"]):
            code = str(row.get("speciesID") or "").strip().lower()
            try:
                gravity = Decimal(str(row.get("spec_gravity")))
                if gravity.is_finite() and gravity > 0:
                    densities[code] = canonical_decimal_string(gravity * 1000)
            except (ValueError, ArithmeticError):
                pass
        accepted = []
        required = {
            "CompetitorID",
            "Event",
            "Time (seconds)",
            "Size (mm)",
            "Species Code",
            "Date (optional)",
        }
        result_records = records(sheets["results"])
        if result_records and not required <= set(result_records[0][1]):
            raise ValueError("legacy Results columns are unsupported")
        for number, row in result_records:
            local_id = str(row.get("CompetitorID") or "").strip()
            when = row.get("Date (optional)")
            if not isinstance(when, datetime):
                excluded["undated"] += 1
                continue
            when = (
                when.replace(tzinfo=timezone.utc)
                if when.tzinfo is None
                else when.astimezone(timezone.utc)
            )
            if when >= cutoff:
                excluded["at_or_after_cutoff"] += 1
                continue
            if local_id not in competitors:
                excluded["unknown_competitor_id"] += 1
                continue
            event = EVENT_CODES.get(str(row.get("Event") or "").strip().upper())
            if event is None:
                excluded["unsupported_event"] += 1
                continue
            try:
                seconds = Decimal(str(row["Time (seconds)"]))
                size = Decimal(str(row["Size (mm)"]))
                if (
                    not seconds.is_finite()
                    or not 0 < seconds <= 600
                    or not size.is_finite()
                    or size <= 0
                    or size != int(size)
                ):
                    raise ValueError
                milliseconds = int(
                    (seconds * 1000).quantize(Decimal("1"), rounding=ROUND_HALF_EVEN)
                )
                if milliseconds <= 0:
                    raise ValueError
            except (ValueError, ArithmeticError):
                excluded["invalid_numeric_fact"] += 1
                continue
            species = str(row.get("Species Code") or "").strip().lower()
            properties = (
                (ContextProperty("density", densities[species], "kg_m3", None),)
                if species in densities
                else ()
            )
            context = TargetContext(
                event, int(size), species, "strathex:v1", "strathex:v1", properties
            )
            note = str(row.get("Notes (Competition, special circumstances, etc.)") or "").strip()
            # A recorded annual competition label is grouped as a whole. Missing
            # labels use the whole day; none of these are claimed live identities.
            group = (
                {"year": when.year, "legacy_label": note}
                if note
                else {"date": when.date().isoformat()}
            )
            tournament = deterministic_identifier("tournament", {"legacy_cohort": group})
            identity = {"workbook_sha256": digest, "sheet": "Results", "row": number}
            accepted.append((when, number, local_id, context, tournament, milliseconds, identity))
        observations = []
        for sequence, (
            when,
            number,
            local_id,
            context,
            tournament,
            milliseconds,
            identity,
        ) in enumerate(sorted(accepted, key=lambda item: (item[0], item[1])), 1):
            observations.append(
                ResultObservation(
                    evidence_id=deterministic_identifier("evidence", identity),
                    competitor_id=competitors[local_id],
                    tournament_id=tournament,
                    round_id=deterministic_identifier("round", {"legacy_cohort": str(tournament)}),
                    field_id=deterministic_identifier("field", identity),
                    context=context,
                    observation_sequence=sequence,
                    occurred_at_utc=when.isoformat(timespec="milliseconds").replace("+00:00", "Z"),
                    # Legacy raw completions have no issued-mark fact. This required
                    # transport value is never used as historical ability evidence.
                    issued_mark=3,
                    completion_clock_ms=None,
                    placing=None,
                    gap_ms=None,
                    result=OfficialResult(ResultStatus.COMPLETION, milliseconds, None, 1, None),
                    source_digest=canonical_digest(
                        {
                            **identity,
                            "raw_time_ms": milliseconds,
                            "context": context.to_dict(),
                            "occurred_at_utc": when.isoformat(),
                        }
                    ),
                )
            )
    finally:
        workbook.close()
    if sha256(path.read_bytes()).hexdigest() != digest:
        raise ValueError("workbook changed while creating the historical snapshot")
    return WorkbookHistory(
        digest,
        tuple(observations),
        tuple(sorted((key, str(value)) for key, value in competitors.items())),
        tuple(sorted(excluded.items())),
    )
