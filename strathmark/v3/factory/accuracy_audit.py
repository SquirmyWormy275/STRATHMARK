"""Private accuracy triage, fixed regression gates and prospective model freezing.

Scores are recomputed from all row receipts. These development checks grant no
promotion, issue authority, Windows qualification or independent blind audit.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from statistics import median

from strathmark.v3.contracts.canonical import canonical_digest
from strathmark.v3.contracts.forecasts import PositiveTimeDistribution
from strathmark.v3.domain.credibility import _quantile_crps
from strathmark.v3.infrastructure.integrity import (
    IntegrityKeyIdentity,
    IntegrityTrustStore,
    P256EphemeralSigner,
    SignedManifest,
    sign_manifest,
    verify_manifest,
)

POLICY = {
    "schema_version": "strathmark-accuracy-regression-policy-v1",
    "minimum_gate_slice_rows": 20,
    "minimum_prospective_rows": 100,
    "minimum_prospective_tournaments": 10,
    "minimum_90_coverage": 0.85,
    "maximum_90_coverage": 0.98,
    "mae_relative_tolerance": 0.0,
    "mae_absolute_tolerance_seconds": 0.001,
    "p90_relative_tolerance": 0.10,
    "p99_relative_tolerance": 0.15,
    "tail_absolute_tolerance_seconds": 2.0,
    "maximum_error_absolute_tolerance_seconds": 5.0,
    "crps_relative_tolerance": 0.05,
    "crps_absolute_tolerance_seconds": 0.10,
    "slice_mae_relative_tolerance": 0.10,
    "slice_mae_absolute_tolerance_seconds": 2.0,
}
COMPONENTS = ("bundle_digest", "formula_digest", "source_implementation_digest")


def attest_benchmark(summary, rows, receipt_file_sha256, *, now=None):
    """Bind local development receipts; this is not independently trusted execution."""
    signer = P256EphemeralSigner.generate("integrity-key:development-benchmark")
    value = {
        **summary,
        "row_receipts_digest": canonical_digest(rows, max_bytes=50_000_000, max_items=1_000_000),
        "row_receipts_sha256": receipt_file_sha256,
        "benchmark_identity": signer.identity.to_dict(),
    }
    now = now or datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace(
        "+00:00", "Z"
    )
    value["benchmark_attestation"] = sign_manifest(
        "development_benchmark",
        {"summary_digest": canonical_digest(value)},
        signer=signer,
        created_at=now,
    ).to_dict()
    return value


def verify_benchmark(summary, rows, *, receipt_file_sha256=None):
    if canonical_digest(rows, max_bytes=50_000_000, max_items=1_000_000) != summary.get(
        "row_receipts_digest"
    ):
        raise ValueError("benchmark row receipts differ from the attested summary")
    if receipt_file_sha256 is not None and receipt_file_sha256 != summary.get(
        "row_receipts_sha256"
    ):
        raise ValueError("benchmark row-file checksum differs from the attested summary")
    value = {k: v for k, v in summary.items() if k != "benchmark_attestation"}
    manifest = SignedManifest.from_dict(summary["benchmark_attestation"])
    identity = IntegrityKeyIdentity.from_dict(summary["benchmark_identity"])
    payload = verify_manifest(manifest, IntegrityTrustStore((identity,)))
    if manifest.kind != "development_benchmark" or payload != {
        "summary_digest": canonical_digest(value)
    }:
        raise ValueError("benchmark summary differs from its development attestation")


def _utc(value):
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        raise ValueError("accuracy timestamps require an explicit timezone")
    return parsed.astimezone(timezone.utc)


def history_band(depth):
    return (
        "none" if depth == 0 else "1_to_4" if depth < 5 else "5_to_14" if depth < 15 else "15_plus"
    )


def _validated(rows):
    if not isinstance(rows, (list, tuple)) or not rows:
        raise ValueError("accuracy needs a nonempty complete row cohort")
    seen = set()
    for row in rows:
        identity = row.get("row_id")
        if not isinstance(identity, str) or not identity or identity in seen:
            raise ValueError("accuracy row identities must be unique")
        seen.add(identity)
        for key in ("tournament_id", "event", "species"):
            if not isinstance(row.get(key), str) or not row[key]:
                raise ValueError("accuracy row context is incomplete")
        for key in ("actual_seconds", "pool_seconds"):
            value = row.get(key)
            if (
                isinstance(value, bool)
                or not isinstance(value, (float, int))
                or not math.isfinite(value)
                or not 0 < value <= 600
            ):
                raise ValueError("accuracy raw seconds must be finite and within declared support")
        for key in ("size_mm", "history_depth"):
            value = row.get(key)
            if (
                isinstance(value, bool)
                or not isinstance(value, int)
                or value < (1 if key == "size_mm" else 0)
            ):
                raise ValueError("accuracy diameter/history depth is invalid")
        distribution = PositiveTimeDistribution.from_dict(row["pool"])
        probabilities = {q.probability for q in distribution.quantiles}
        if not {"0.05", "0.25", "0.5", "0.75", "0.95"} <= probabilities:
            raise ValueError("accuracy needs exact common quantiles, including 90% endpoints")
        if abs(distribution.median_ms / 1000 - row["pool_seconds"]) > 0.001000001:
            raise ValueError("accuracy median differs from the supplied distribution")
    return rows


def _metrics(rows):
    errors, coverage, widths, scores = [], [], [], []
    for row in rows:
        errors.append(abs(row["pool_seconds"] - row["actual_seconds"]))
        distribution = PositiveTimeDistribution.from_dict(row["pool"])
        quantiles = {q.probability: q.time_ms / 1000 for q in distribution.quantiles}
        coverage.append(quantiles["0.05"] <= row["actual_seconds"] <= quantiles["0.95"])
        widths.append(quantiles["0.95"] - quantiles["0.05"])
        scores.append(
            float(_quantile_crps(distribution, Decimal(str(row["actual_seconds"] * 1000)))) / 1000
        )
    ordered = sorted(errors)
    return {
        "row_count": len(rows),
        "mean_absolute_error_seconds": sum(errors) / len(rows),
        "median_absolute_error_seconds": median(errors),
        "p90_absolute_error_seconds": ordered[math.ceil(0.90 * len(rows)) - 1],
        "p99_absolute_error_seconds": ordered[math.ceil(0.99 * len(rows)) - 1],
        "maximum_absolute_error_seconds": ordered[-1],
        "interval_90_coverage": sum(coverage) / len(rows),
        "mean_90_interval_width_seconds": sum(widths) / len(rows),
        "mean_quantile_crps_seconds": sum(scores) / len(rows),
    }


def triage_rows(rows):
    rows = _validated(rows)
    categories = defaultdict(lambda: defaultdict(list))
    for row in rows:
        for key in ("event", "species", "size_mm"):
            categories[key][str(row[key])].append(row)
        categories["history_depth"][history_band(row["history_depth"])].append(row)
        if "same_event_history_depth" in row and "same_material_history_depth" in row:
            event_depth, material_depth = (
                row["same_event_history_depth"],
                row["same_material_history_depth"],
            )
            if (
                any(
                    isinstance(v, bool) or not isinstance(v, int) or v < 0
                    for v in (event_depth, material_depth)
                )
                or not material_depth <= event_depth <= row["history_depth"]
            ):
                raise ValueError("relevant-history support counts are inconsistent")
            band = (
                "no_event_history"
                if not event_depth
                else "event_without_material_history"
                if not material_depth
                else "material_history"
            )
            categories["relevant_history"][band].append(row)
    ranked = sorted(rows, key=lambda r: abs(r["pool_seconds"] - r["actual_seconds"]), reverse=True)
    count = math.ceil(0.1 * len(rows))
    total = sum(abs(r["pool_seconds"] - r["actual_seconds"]) for r in rows)
    worst = sum(abs(r["pool_seconds"] - r["actual_seconds"]) for r in ranked[:count])
    return {
        "schema_version": "strathmark-accuracy-triage-v1",
        "overall": _metrics(rows),
        "tournament_count": len({r["tournament_id"] for r in rows}),
        "tail": {
            "worst_decile_count": count,
            "worst_decile_share_of_total_error": worst / total if total else 0.0,
        },
        "slices": {
            k: {v: _metrics(group) for v, group in sorted(groups.items())}
            for k, groups in sorted(categories.items())
        },
        "all_rows_retained": True,
        "numeric_promotion": False,
    }


def compare_rows(baseline_rows, candidate_rows):
    baseline_rows, candidate_rows = _validated(baseline_rows), _validated(candidate_rows)
    old = {r["row_id"]: r for r in baseline_rows}
    new = {r["row_id"]: r for r in candidate_rows}
    if old.keys() != new.keys():
        raise ValueError("accuracy comparison must retain the identical complete cohort")
    for identity, row in old.items():
        other = new[identity]
        if any(
            row[k] != other[k]
            for k in ("tournament_id", "event", "species", "size_mm", "history_depth")
        ) or round(row["actual_seconds"] * 1000) != round(other["actual_seconds"] * 1000):
            raise ValueError("accuracy comparison changed a target, context or historical support")
        for key in ("same_event_history_depth", "same_material_history_depth", "occurred_at_utc"):
            if (key in row) != (key in other) or (key in row and row[key] != other[key]):
                raise ValueError("accuracy comparison changed causal evidence context")
    baseline, candidate = triage_rows(baseline_rows), triage_rows(candidate_rows)
    violations = []

    def upper(metric, relative, absolute, scope="overall", before=None, after=None):
        before = baseline["overall"] if before is None else before
        after = candidate["overall"] if after is None else after
        limit = before[metric] * (1 + relative) + absolute
        if after[metric] > limit + 1e-9:
            violations.append(
                {"scope": scope, "metric": metric, "value": after[metric], "maximum": limit}
            )

    upper(
        "mean_absolute_error_seconds",
        POLICY["mae_relative_tolerance"],
        POLICY["mae_absolute_tolerance_seconds"],
    )
    upper(
        "p90_absolute_error_seconds",
        POLICY["p90_relative_tolerance"],
        POLICY["tail_absolute_tolerance_seconds"],
    )
    upper(
        "p99_absolute_error_seconds",
        POLICY["p99_relative_tolerance"],
        POLICY["tail_absolute_tolerance_seconds"],
    )
    upper("maximum_absolute_error_seconds", 0, POLICY["maximum_error_absolute_tolerance_seconds"])
    upper(
        "mean_quantile_crps_seconds",
        POLICY["crps_relative_tolerance"],
        POLICY["crps_absolute_tolerance_seconds"],
    )
    coverage = candidate["overall"]["interval_90_coverage"]
    if not POLICY["minimum_90_coverage"] <= coverage <= POLICY["maximum_90_coverage"]:
        violations.append(
            {
                "scope": "overall",
                "metric": "interval_90_coverage",
                "value": coverage,
                "minimum": POLICY["minimum_90_coverage"],
                "maximum": POLICY["maximum_90_coverage"],
            }
        )
    for category, groups in baseline["slices"].items():
        for label, before in groups.items():
            if before["row_count"] < POLICY["minimum_gate_slice_rows"]:
                continue
            if label not in candidate["slices"].get(category, {}):
                raise ValueError("candidate omitted an eligible accuracy slice")
            upper(
                "mean_absolute_error_seconds",
                POLICY["slice_mae_relative_tolerance"],
                POLICY["slice_mae_absolute_tolerance_seconds"],
                f"{category}:{label}",
                before,
                candidate["slices"][category][label],
            )
    return {
        "schema_version": "strathmark-accuracy-regression-report-v1",
        "policy": POLICY,
        "baseline": baseline,
        "candidate": candidate,
        "accepted": not violations,
        "violations": violations,
        "cohort_count": len(old),
        "numeric_promotion": False,
        "qualification": "development regression checks only",
    }


def freeze_protocol(summary, examined_rows, *, now=None):
    rows = _validated(examined_rows)
    verify_benchmark(summary, rows)
    now = now or datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace(
        "+00:00", "Z"
    )
    frozen_at = _utc(now)
    if any(_utc(r["occurred_at_utc"]) > frozen_at for r in rows):
        raise ValueError("protocol freeze cannot precede examined results")
    for key in COMPONENTS:
        value = summary.get(key)
        if (
            not isinstance(value, str)
            or len(value) != 64
            or any(c not in "0123456789abcdef" for c in value)
        ):
            raise ValueError("prospective freezing requires exact component digests")
    value = {
        "schema_version": "strathmark-prospective-accuracy-protocol-v1",
        "created_at_utc": now,
        "not_before_utc": now,
        "components": {k: summary[k] for k in COMPONENTS},
        "examined_receipts_digest": summary["row_receipts_digest"],
        "examined_benchmark_attestation_digest": SignedManifest.from_dict(
            summary["benchmark_attestation"]
        ).body_digest,
        "examined_row_ids": sorted(r["row_id"] for r in rows),
        "examined_tournament_ids": sorted({r["tournament_id"] for r in rows}),
        "regression_policy": POLICY,
        "status": "awaiting_future_competitions",
        "numeric_promotion": False,
        "qualification": "frozen-model prospective development evaluation; no independent blind audit",
    }
    value["digest"] = canonical_digest(value)
    return value


def verify_prospective(protocol, summary, rows):
    verify_benchmark(summary, rows)
    value = {k: v for k, v in protocol.items() if k != "digest"}
    if (
        protocol.get("schema_version") != "strathmark-prospective-accuracy-protocol-v1"
        or canonical_digest(value) != protocol.get("digest")
        or protocol.get("regression_policy") != POLICY
    ):
        raise ValueError("prospective protocol digest or fixed policy changed")
    if any(summary.get(k) != protocol["components"][k] for k in COMPONENTS):
        raise ValueError("prospective benchmark differs from a frozen component")
    rows = _validated(rows)
    if any(
        r["row_id"] in protocol["examined_row_ids"]
        or r["tournament_id"] in protocol["examined_tournament_ids"]
        for r in rows
    ):
        raise ValueError("prospective cohort includes previously examined evidence")
    if any(_utc(r["occurred_at_utc"]) <= _utc(protocol["not_before_utc"]) for r in rows):
        raise ValueError("prospective targets must occur strictly after protocol freeze")
    return {
        "status": "eligible_for_frozen_model_evaluation",
        "row_count": len(rows),
        "tournament_count": len({r["tournament_id"] for r in rows}),
        "protocol_digest": protocol["digest"],
        "numeric_promotion": False,
    }


def audit_workbook(path, *, cutoff_at_utc):
    """Report import quality and possible repeated rows without editing evidence."""
    from collections import Counter

    from openpyxl import load_workbook

    from strathmark.v3.factory.workbook_history import load_workbook_history

    path = Path(path).resolve(strict=True)
    history = load_workbook_history(path, cutoff_at_utc=cutoff_at_utc)
    workbook = load_workbook(path, read_only=True, data_only=True)
    repeated, first, labels, count = [], {}, Counter(), 0
    try:
        sheet = next(s for s in workbook if s.title.lower() == "results")
        values = iter(sheet.values)
        headers = next(values)
        if len([h for h in headers if h is not None]) != len(
            set(h for h in headers if h is not None)
        ):
            raise ValueError("Results repeats a column label")
        for number, values in enumerate(values, 2):
            count += 1
            row = dict(zip(headers, values))
            when = row.get("Date (optional)")
            if isinstance(when, datetime):
                key = tuple(
                    str(row.get(k))
                    for k in (
                        "CompetitorID",
                        "Event",
                        "Size (mm)",
                        "Species Code",
                        "Time (seconds)",
                        "Date (optional)",
                    )
                )
                if key in first:
                    repeated.append([first[key], number])
                else:
                    first[key] = number
            note = str(row.get("Notes (Competition, special circumstances, etc.)") or "").strip()
            labels["recorded_label" if note else "whole_date_fallback"] += 1
    finally:
        workbook.close()
    if hashlib.sha256(path.read_bytes()).hexdigest() != history.source_sha256:
        raise ValueError("workbook changed during read-only quality audit")
    return {
        "schema_version": "strathmark-workbook-quality-audit-v1",
        "workbook_sha256": history.source_sha256,
        "source_result_rows": count,
        "dated_admitted_completions": len(history.observations),
        "excluded": dict(history.excluded),
        "possible_repeated_row_pairs": repeated,
        "possible_repeated_row_count": len(repeated),
        "grouping": dict(labels),
        "latest_admitted_result_utc": max(
            (o.occurred_at_utc for o in history.observations), default=None
        ),
        "live_data_modified": False,
        "rows_deleted": 0,
        "interpretation": "repeated rows require field/round confirmation; equal raw cuts can be legitimate separate heats",
        "time_definition": "declared raw cutting seconds; no issued mark is subtracted or inferred",
        "date_precision": "legacy calendar timestamps; same-date results never enter an earlier forecast",
    }


def _load(path):
    if path.stat().st_size > 50_000_000:
        raise ValueError("accuracy input exceeds its 50 MB bound")
    return json.loads(path.read_bytes())


def _write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    raw = json.dumps(value, indent=2, allow_nan=False).encode()
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(descriptor, "wb") as stream:
        stream.write(raw)
        stream.flush()
        os.fsync(stream.fileno())


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="operation", required=True)
    for name in ("triage", "freeze", "prospective", "compare"):
        command = sub.add_parser(name)
        command.add_argument("--benchmark", type=Path, required=True)
        command.add_argument("--output", type=Path, required=True)
        if name == "compare":
            command.add_argument("--baseline", type=Path, required=True)
        if name == "prospective":
            command.add_argument("--protocol", type=Path, required=True)
    workbook = sub.add_parser("workbook")
    workbook.add_argument("--workbook", type=Path, required=True)
    workbook.add_argument("--cutoff-at-utc", required=True)
    workbook.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.operation == "workbook":
        result = audit_workbook(args.workbook, cutoff_at_utc=args.cutoff_at_utc)
        _write(args.output, result)
        print(
            json.dumps(
                {
                    k: result[k]
                    for k in (
                        "schema_version",
                        "source_result_rows",
                        "dated_admitted_completions",
                        "excluded",
                        "possible_repeated_row_count",
                        "live_data_modified",
                    )
                },
                indent=2,
            )
        )
        return 0
    rows_path = args.benchmark / "private-row-receipts.json"
    rows, summary = _load(rows_path), _load(args.benchmark / "summary.json")
    verify_benchmark(
        summary, rows, receipt_file_sha256=hashlib.sha256(rows_path.read_bytes()).hexdigest()
    )
    if args.operation == "triage":
        result = triage_rows(rows)
    elif args.operation == "freeze":
        result = freeze_protocol(summary, rows)
    elif args.operation == "prospective":
        result = verify_prospective(_load(args.protocol), summary, rows)
        result["enough_evidence"] = (
            result["row_count"] >= POLICY["minimum_prospective_rows"]
            and result["tournament_count"] >= POLICY["minimum_prospective_tournaments"]
        )
        result["triage"] = triage_rows(rows)
    else:
        result = compare_rows(_load(args.baseline / "private-row-receipts.json"), rows)
    result["receipt_file_sha256"] = hashlib.sha256(rows_path.read_bytes()).hexdigest()
    _write(args.output, result)
    print(
        json.dumps(
            {
                k: result[k]
                for k in (
                    "schema_version",
                    "accepted",
                    "status",
                    "cohort_count",
                    "numeric_promotion",
                    "enough_evidence",
                )
                if k in result
            },
            indent=2,
        )
    )
    return 2 if result.get("accepted") is False or result.get("enough_evidence") is False else 0


if __name__ == "__main__":
    raise SystemExit(main())
